"""One transaction owns reception revisions and conversion to a contract matter."""
from io import BytesIO
import json
from uuid import UUID, uuid4

from sqlalchemy import text

from .. import domain as d
from ..persistence import SessionRow
from ..security import ServiceError, _principal, authorize
from .. import matters
from ..record_codec import encode, now, sha


def _lawyer(actor, db):
    current = _principal(db.get(SessionRow, str(actor.session_id), populate_existing=True), db, user_id=actor.user_id)
    if current is None:
        raise ServiceError('AUTHORIZATION_REQUIRED', 401)
    if not current.roles & {d.Role.HANDLING_LAWYER, d.Role.REVIEWING_LAWYER}:
        raise ServiceError('CROSS_MATTER_DENIED', 403)
    return current


def _owned(identifier, actor, db):
    _lawyer(actor, db)
    row = db.execute(text('SELECT * FROM intake_cases WHERE id=:id AND owner_id=:owner'),
        {'id':str(identifier), 'owner':str(actor.user_id)}).mappings().first()
    if row is None:
        raise ServiceError('CROSS_MATTER_DENIED', 403)
    if row['matter_id']:
        authorize(actor, UUID(row['matter_id']), 'read', db)
    return row


def _consultation_source(db, identifier, revision):
    if not db.scalar(text("SELECT 1 FROM sqlite_master WHERE type='table' AND name='assistant_profile_receipts'")):
        return {}
    source = db.execute(text('''SELECT consultation_origin, conversation_id, profile_revision, intake_revision
        FROM assistant_profile_receipts WHERE intake_id=:id AND intake_revision<=:revision
        AND projection_changed=1 ORDER BY intake_revision DESC LIMIT 1'''),
        {'id':str(identifier), 'revision':revision}).mappings().first()
    if source is None:
        return {}
    rows = db.execute(text('''SELECT revision,payload FROM intake_revisions WHERE intake_id=:id
        AND revision>=:first AND revision<=:last ORDER BY revision'''),
        {'id':str(identifier), 'first':source['intake_revision'], 'last':revision}).mappings().all()
    # A manual replacement breaks provenance even if a later edit restores identical text.
    if len(rows) != revision-source['intake_revision']+1 or len({json.loads(r['payload'])['consultation'] for r in rows}) != 1:
        return {}
    return {'consultation_origin':source['consultation_origin'],
            'source_conversation_id':source['conversation_id'],
            'source_profile_revision':source['profile_revision'] or 0}


def read(identifier, actor, db, revision=None):
    owner = _owned(identifier, actor, db)
    row = db.execute(text('SELECT * FROM intake_revisions WHERE intake_id=:id AND revision=:revision'),
        {'id':str(identifier), 'revision':revision or owner['current_revision']}).mappings().first()
    if row is None:
        raise ServiceError('VERSION_CONFLICT', 409)
    metadata = _consultation_source(db, identifier, row['revision'])
    return d.IntakeRecord(id=identifier, revision=row['revision'], saved_at=row['saved_at'],
        matter_id=owner['matter_id'], client_reply=row['client_reply'], internal_notes=row['internal_notes'],
        **json.loads(row['payload']), **metadata)


def listing(actor, db, offset=0):
    if not actor.roles & {d.Role.HANDLING_LAWYER, d.Role.REVIEWING_LAWYER}:
        return {'items':[], 'next_offset':None}
    _lawyer(actor, db)
    # Revoked converted cases disappear from the list as well as direct routes.
    rows = db.execute(text('''SELECT c.id FROM intake_cases c WHERE c.owner_id=:owner AND
        (c.matter_id IS NULL OR EXISTS (SELECT 1 FROM matter_memberships m WHERE m.matter_id=c.matter_id AND m.user_id=:owner))
        ORDER BY c.rowid DESC LIMIT 51 OFFSET :offset'''), {'owner':str(actor.user_id), 'offset':offset}).all()
    items=[]
    for row in rows[:50]:
        record=read(UUID(row[0]),actor,db)
        items.append({'id':str(record.id),'title':record.title,'revision':record.revision,
            'status':'converted' if record.matter_id else record.status,
            'follow_up_at':record.follow_up_at.isoformat() if record.follow_up_at else None,
            'follow_up_note':record.follow_up_note, 'materials_needed':sum(m.state=='needed' for m in record.materials),
            'saved_at':record.saved_at.isoformat(), 'matter_id':str(record.matter_id) if record.matter_id else None})
    return {'items':items, 'next_offset':offset+50 if len(rows)>50 else None}


def _snapshot(identifier, revision, value, db):
    db.execute(text('''INSERT INTO intake_revisions
        (intake_id,revision,saved_at,payload,client_reply,internal_notes)
        VALUES (:id,:revision,:at,:payload,:client,:internal)'''),
        {'id':str(identifier),'revision':revision,'at':now().isoformat(),
         'payload':encode(value.model_dump(mode='json',exclude={'client_reply','internal_notes'})),
         'client':value.client_reply,'internal':value.internal_notes})


def create(value, actor, db):
    _lawyer(actor, db)
    digest=sha(value.draft.model_dump(mode='json'))
    previous=db.execute(text('SELECT id,request_sha256 FROM intake_cases WHERE owner_id=:owner AND request_id=:request'),
        {'owner':str(actor.user_id),'request':str(value.request_id)}).mappings().first()
    if previous:
        if previous['request_sha256']!=digest:
            raise ServiceError('VERSION_CONFLICT',409)
        return read(UUID(previous['id']),actor,db)
    identifier=uuid4()
    db.execute(text('''INSERT INTO intake_cases (id,owner_id,request_id,request_sha256,current_revision)
        VALUES (:id,:owner,:request,:digest,1)'''),
        {'id':str(identifier),'owner':str(actor.user_id),'request':str(value.request_id),'digest':digest})
    _snapshot(identifier,1,value.draft,db)
    return read(identifier,actor,db)


def save(identifier, value, actor, db):
    row=_owned(identifier,actor,db)
    if row['matter_id'] or row['current_revision']!=value.expected_revision:
        raise ServiceError('VERSION_CONFLICT',409)
    revision=value.expected_revision+1
    _snapshot(identifier,revision,value.draft,db)
    db.execute(text('UPDATE intake_cases SET current_revision=:revision WHERE id=:id'),
        {'id':str(identifier),'revision':revision})
    return read(identifier,actor,db)


def convert(identifier,value,actor,db):
    row=_owned(identifier,actor,db)
    digest=sha(value.model_dump(mode='json'))
    if row['matter_id']:
        if row['conversion_sha256']!=digest:
            raise ServiceError('VERSION_CONFLICT',409)
        return {'matter_id':row['matter_id'],'intake_id':str(identifier),'revision':row['current_revision']}
    record=read(identifier,actor,db)
    if record.revision!=value.expected_revision or record.status=='closed':
        raise ServiceError('VERSION_CONFLICT',409)
    # The caller's transaction covers both existing service writes and the link.
    matter=matters.create_matter(value.matter,actor,db)
    db.execute(text('UPDATE intake_cases SET matter_id=:matter,conversion_sha256=:digest WHERE id=:id'),
        {'id':str(identifier),'matter':str(matter.id),'digest':digest})
    return {'matter_id':str(matter.id),'intake_id':str(identifier),'revision':record.revision}


def export(identifier,value,actor,db):
    from docx import Document

    record=read(identifier,actor,db)
    if record.revision!=value.revision or (value.audience=='client' and not value.text_checked):
        raise ServiceError('VERSION_CONFLICT',409)
    doc=Document()
    doc.core_properties.author='群兴法律工作台'
    doc.core_properties.title='客户回复草稿' if value.audience=='client' else '内部接待记录草稿'
    doc.add_heading(doc.core_properties.title,0)
    doc.add_paragraph(f'接待编号 {identifier} · 第 {record.revision} 版')
    doc.add_paragraph('人工工作稿，依据仍待核验；导出不等于正式签发或发送。')
    if value.audience=='client':
        if not record.client_reply.strip():
            raise ServiceError('STATE_CONFLICT',409)
        # Deliberate allowlist: no draft title, contact, history, or internal text.
        doc.add_paragraph(record.client_reply)
    else:
        if record.source_conversation_id is not None:
            doc.add_paragraph(f'来源会话 {record.source_conversation_id} · 档案第 {record.source_profile_revision} 版 · 待核验')
        for heading,content in (
            ('咨询主题',record.title),('联系备注',record.contact),
            ('接待摘要' if record.consultation_origin=='profile_summary' else '咨询原话',record.consultation),
            ('追问清单','\n'.join(record.questions)),('材料清单','\n'.join(f'{m.name}（'+
             {'needed':'待补','received':'已收到，待核','checked':'已核对'}[m.state]+'）' for m in record.materials)),
            ('跟进', (record.follow_up_at.isoformat()+' UTC\n' if record.follow_up_at else '')+record.follow_up_note),
            ('律师内部意见',record.internal_notes)):
            doc.add_heading(heading,2);doc.add_paragraph(content or '尚未填写')
    stream=BytesIO();doc.save(stream)
    return stream.getvalue()
