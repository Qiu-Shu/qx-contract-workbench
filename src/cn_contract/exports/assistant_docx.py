"""Fresh Word package containing only the already filtered output projection."""
from io import BytesIO


def build(preview):
    # Keep Word's dependency out of normal chat startup, as in the intake exporter.
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ('Normal','Title','Heading 1','Header','Footer'):
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
        for border in list(style.element.findall('.//'+qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.4
    document.styles['Title'].font.size = Pt(20)
    document.styles['Heading 1'].font.size = Pt(13)
    section.header.paragraphs[0].text = preview.title+' · 待律师核对'
    section.footer.paragraphs[0].text = '工作草稿 · 非正式法律意见'
    document.add_paragraph(preview.title,'Title')
    document.add_paragraph(preview.source_label)
    for note in preview.notes:
        document.add_paragraph(note)
    document.add_heading('回复正文' if preview.selection.audience == 'client' else '工作稿正文',1)
    for line in preview.text.split('\n'):
        document.add_paragraph(line)
    if preview.citations:
        document.add_heading('来源核对',1)
        for citation in preview.citations:
            document.add_paragraph(citation)
    core = document.core_properties
    core.title = preview.title
    core.subject = preview.source_label
    core.author = core.last_modified_by = '群兴法律工作台'
    core.comments = '待律师核对 · 非正式法律意见'
    core.keywords = 'DRAFT '+preview.selection.audience
    output = BytesIO()
    document.save(output)
    return output.getvalue()
