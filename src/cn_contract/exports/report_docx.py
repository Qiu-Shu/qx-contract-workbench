"""Fresh A4 Chinese draft, never a rewrite of an untrusted original package."""
from io import BytesIO
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _table(document,head,rows):
    if not rows:return
    table=document.add_table(rows=1,cols=len(head));table.style='Table Grid'
    for cell,text in zip(table.rows[0].cells,head):cell.text=text
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=table.add_row().cells
        for cell,value in zip(cells,values):cell.text=str(value)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(4)


def build(projection):
    doc=Document()
    watermark=projection.get('watermark','DRAFT MOCK 模拟材料 不构成正式法律意见')
    doc.settings.odd_and_even_pages_header_footer=False
    section=doc.sections[0];section.different_first_page_header_footer=False;section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2);section.left_margin=section.right_margin=Cm(2.2)
    section.header_distance=section.footer_distance=Cm(0.8)
    for name in ('Normal','Title','Heading 1','Heading 2','Header','Footer'):
        style=doc.styles[name]
        for border in list(style.element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
        style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(20)
    doc.styles['Title'].font.underline=False
    doc.styles['Heading 1'].font.size=Pt(14)
    for name in ('Heading 1','Heading 2'):doc.styles[name].paragraph_format.keep_with_next=True
    for container in (section.header,section.footer):
        paragraph=container.paragraphs[0];paragraph.text=watermark
        for run in paragraph.runs:run.font.size=Pt(9)
    doc.add_paragraph(projection['title'],'Title')
    doc.add_paragraph(projection['opening'])
    if projection.get('scope'):
        doc.add_heading('读取材料范围',1)
        for text in projection['scope']:doc.add_paragraph(text)
    if projection.get('internal'):
        doc.add_heading('内部立场、底线与复核任务',1)
        for text in projection['internal']:doc.add_paragraph(text)
    doc.add_heading('使用范围与材料限制',1)
    for text in projection['limitations']:doc.add_paragraph(text)
    if projection['coverage']:
        doc.add_heading('审查维度与处理状态',1)
        _table(doc,('维度','状态'),[(c['dimension'],c['status']) for c in projection['coverage']])
    if projection['observations']:
        doc.add_heading('材料观察与人工复核事项',1)
        for index,item in enumerate(projection['observations'],1):
            doc.add_heading(f'复核事项 {index}',2)
            for key in ('decision_status','analysis','effect','recommendation','basis_review','replacement','fallback','original','history'):
                if key in item:doc.add_paragraph(item[key])
    doc.add_heading(projection.get('patch_heading','已批准的模拟修改文字'),1)
    _table(doc,('原文','建议新文','修改说明'),[(p['old'],p['new'],p['reason']) for p in projection['patches']])
    if not projection['patches']:doc.add_paragraph(projection.get('empty_patches','当前没有经模拟复核身份批准的修改文字。'))
    doc.add_heading('需要人工处理的修改',1)
    _table(doc,('位置与状态','已批准文字及人工处理说明'),[(m['location']+'\n需人工处理',
        '原文：'+m['old']+'\n建议新文：'+m['new']+'\n批准修改说明：'+m['explanation']+'\n人工处理原因：'+m['reason']) for m in projection['manual']])
    if not projection['manual']:doc.add_paragraph('本次所列修改没有未写入的人工项。其他材料及法律判断仍须律师复核。')
    core=doc.core_properties
    core.title=projection['title'];core.subject=projection.get('subject','独立虚构模拟合同草案')
    core.author=projection.get('author','群兴工程模拟');core.last_modified_by=core.author
    core.comments=watermark if 'watermark' in projection else 'DRAFT MOCK 非正式法律意见'
    core.keywords=projection.get('keywords','DRAFT MANUAL-SYNTHETIC' if 'watermark' in projection else 'DRAFT MOCK')
    from datetime import datetime,timezone
    core.created=core.modified=datetime.now(timezone.utc)
    output=BytesIO();doc.save(output);return output.getvalue()
