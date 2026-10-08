from pathlib import Path
from hashlib import sha256
from datetime import datetime,timezone
import json,subprocess,shutil,re
from docx import Document
from pypdf import PdfReader
ROOT=Path.cwd();OUT=ROOT/'.audit-local/word-layout-20261009'
DEST=ROOT/'docs/testing/ocr-word-layout-20261009'
DEST.mkdir(exist_ok=True)
jobs={'final':'20261008T205410Z-c32c37e7-0505-4324-991c-07ddd9442bab','manual':'20261008T205422Z-c5c17718-6035-4216-a105-16f28bfd79c7'}
records=[]
normalize=lambda value:re.sub(r'\s+','',value)
for label,name in jobs.items():
    work=ROOT/'.audit-local/t09-word-qa'/name
    driver=json.loads((work/'driver.json').read_text(encoding='utf8'))
    result=json.loads((work/'result.json').read_text(encoding='utf-8-sig'))
    pages=json.loads((work/'pages.json').read_text(encoding='utf8'))
    markers=json.loads((work/'png-markers.json').read_text(encoding='utf8'))
    source=Path(driver['original_path']);raw=source.read_bytes()
    assert sha256(raw).hexdigest()==driver['original_sha256']==result['input_sha256']
    assert raw==subprocess.check_output(['git','show','HEAD:'+source.relative_to(ROOT).as_posix()])
    assert driver['exit_code']==0 and result['status']=='WORD_EXPORT_EXECUTED'
    assert result['macros_disabled'] and result['window_pid_verified'] and not result['repair_requested'] and not result['add_recent']
    assert result['revisions_before']==result['revisions_after']==0
    pdf=PdfReader(work/'output.pdf');assert len(pdf.pages)==pages['pages']==4
    pdftext=normalize('\n'.join(p.extract_text() for p in pdf.pages))
    doc=Document(source)
    paragraphs=[p.text for p in doc.paragraphs if p.text.strip()]
    cells=[c.text for table in doc.tables for row in table.rows for c in row.cells if c.text.strip()]
    missing=[t for t in paragraphs+cells if normalize(t) not in pdftext]
    assert not missing,missing
    for page in pdf.pages:
        positions=[]
        def visitor(t,cm,tm,font,size):
            if 'DRAFT' in t:positions.append(tm[5])
        page.extract_text(visitor_text=visitor)
        assert any(y>float(page.mediabox.height)-60 for y in positions)
        assert any(y<60 for y in positions)
    assert markers['all_pages_pixel_identical_markers']
    target=DEST/label;target.mkdir(exist_ok=True)
    for fname in ['driver.json','result.json','pages.json','png-markers.json','output.pdf','pdfinfo.log','pdftoppm.log']+[f'page-{n}.png' for n in range(1,5)]:
        shutil.copyfile(work/fname,target/fname)
    for item in pages['images']:
        assert sha256(Path(item['path']).read_bytes()).hexdigest()==item['sha256']
    records.append({'file':source.relative_to(ROOT).as_posix(),'sha256':driver['original_sha256'],'pages':4,'paragraphs_preserved':len(paragraphs),'table_cells_preserved':len(cells),'word_version':result['word_version'],'word_build':result['word_build'],'original_unchanged':True,'header_footer_all_pages':True,'visual_review':'All four page PNGs opened and inspected; no clipping, overlap, missing Chinese glyphs or broken table observed. Dense diagnostic codes remain a readability limitation.'})
receipts=[]
for p in OUT.glob('*.json'):
    value=json.loads(p.read_text(encoding='utf8'))
    if isinstance(value,dict) and 'stage' in value:
        receipts.append(value);shutil.copyfile(p,DEST/p.name)
        if p.with_suffix('.txt').exists():shutil.copyfile(p.with_suffix('.txt'),DEST/p.with_suffix('.txt').name)
record={'checked_utc':datetime.now(timezone.utc).isoformat(),'status':'NATIVE_WORD_TWO_FICTIONAL_DRAFTS_EIGHT_PAGES_VISUALLY_CHECKED','records':records,'pages_total':8,'source_docx_edits':0,'product_code_changes':0,'new_packages':0,'real_model_or_qcc_calls':0,'native_renderer':'installed Microsoft Word; existing T09 isolated read-only helper','canonical_renderer':'BLOCKED_MISSING_SOFFICE','receipts':sorted(receipts,key=lambda x:x['started_utc'])}
(DEST/'verification.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps({'status':record['status'],'records':records},ensure_ascii=False))
