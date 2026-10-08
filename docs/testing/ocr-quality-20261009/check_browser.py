from pathlib import Path
import json,hashlib,sqlite3
from docx import Document
from uuid import UUID
from sqlalchemy import select
from cn_contract.config import load_settings
from cn_contract.persistence import Database,UserRow,SessionRow
from cn_contract.security import _principal,ServiceError
from cn_contract.domain import SyntheticAssistedPolicy
from cn_contract.review.proposals import configure
from cn_contract.exports.approvals import download_artifact
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent
aid='08363003-e614-4a6b-8ac2-6353d08d8ea3'
raw=(Path('D:/Downloads')/f'draft-{aid}.docx').read_bytes()
dest=OUT/'browser-final-draft.docx';dest.write_bytes(raw)
doc=Document(dest)
text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
checks=['28473','可见原文核对','非PDF文字层','b17b8fc7-825b-4021-93da-a8a2fb5d94ee','017866f3c56b2fad']
assert all(v in text for v in checks)
(OUT/'browser-final-draft.txt').write_text(text,encoding='utf8')
receipt={'file':f'draft-{aid}.docx','bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest(),'content_checks':checks,'visual_layout':'NOT_VERIFIED'}
(OUT/'browser-final-draft-check.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
root=ROOT/'.local/ocr-quality-preview-20261009'
database=Database('sqlite:///'+str(root/'workbench.db'),root/'protected',load_settings(),initialize_schema=False)
from cn_contract.web.app import initialize_storage
initialize_storage(database)
configure(database,'assisted_draft',SyntheticAssistedPolicy(policy_id='ocr-quality-preview-20261009',purpose='assisted_draft',document_hashes=('017866f3c56b2fad241fc048e6440e30314db41bef53391eabd7bc706117ae9b',),candidate_source='manual_synthetic'))
with database.session() as db:
    user=db.scalar(select(UserRow).where(UserRow.username=='lawyer'))
    session=db.scalar(select(SessionRow).where(SessionRow.user_id==user.id,SessionRow.revoked==False))
    actor=_principal(session,db)
    assert actor is not None
    try:download_artifact(UUID('a5c16694-1d1c-491c-8246-2a96b15cd087'),actor,db=db)
    except ServiceError as e:
        stale={'code':e.code,'status':e.status_code if hasattr(e,'status_code') else e.status}
        assert stale['status']==409
    else:raise AssertionError('stale draft accepted')
    assert download_artifact(UUID(aid),actor,db=db).getvalue()==raw
con=sqlite3.connect(f'file:{root.as_posix()}/workbench.db?mode=ro',uri=True)
tables=[r[0] for r in con.execute("select name from sqlite_master where type='table'")]
counts={t:con.execute('select count(*) from "'+t+'"').fetchone()[0] for t in tables}
sources=[json.loads(r[0]) for r in con.execute('select payload from page_verifications')]
candidates=[json.loads(r[0]) for r in con.execute('select payload from ocr_candidates')]
assert len(candidates)==1 and len(sources)==2
linked=next(s for s in sources if s['input']['candidate_id'])
assert linked['input']['candidate_id']==candidates[0]['result']['id']
assert linked['input']['candidate_sha256']==candidates[0]['sha256']
assert all(b['verification_state']=='unverified' for b in candidates[0]['result']['blocks'])
assert counts['approvals']==counts['review_model_attempts']==counts['query_ledger']==0
records=[json.loads(r[0]) for r in con.execute("select payload from export_records where kind='working_draft_confirmation'")]
evidence={'old_download_service':stale,'new_download_matches_browser':True,'candidate_source_link_verified':True,'candidate_blocks_unverified':True,'counts':counts,'sources':sources,'candidates':candidates,'working_draft_confirmations':records}
(OUT/'browser-db-check.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps({'draft':receipt,'old_download_service':stale,'counts':counts},ensure_ascii=False))
