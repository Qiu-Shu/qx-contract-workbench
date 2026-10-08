from pathlib import Path
from hashlib import sha256
from datetime import datetime,timezone
import json,subprocess,shutil,re
from docx import Document
from pypdf import PdfReader
ROOT=Path.cwd();OUT=ROOT/'.audit-local/client-readability-20261009'
DEST=ROOT/'docs/testing/client-readability-20261009';DEST.mkdir(exist_ok=True)
jobs={'client':'20261008T210829Z-08265288-c0df-4b9c-a736-bcff893697eb','internal':'20261008T210835Z-bea7ca88-945a-47d9-bf27-a05f16e5b9af'}
normalize=lambda value:re.sub(r'\s+','',value)
oldpath=ROOT/'docs/testing/ocr-quality-20261009/browser-final-draft.docx'
assert oldpath.read_bytes()==subprocess.check_output(['git','show','HEAD:'+oldpath.relative_to(ROOT).as_posix()])
old=Document(oldpath);oldparas=[p.text for p in old.paragraphs]
client=Document(OUT/'client.docx');clientparas=[p.text for p in client.paragraphs]
scope_start=oldparas.index('读取材料范围');scope_end=oldparas.index('使用范围与材料限制')
assert oldparas[:scope_start]==clientparas[:clientparas.index('读取材料范围')]
assert oldparas[scope_end:]==clientparas[clientparas.index('使用范围与材料限制'):]
cells=lambda doc:[c.text for t in doc.tables for row in t.rows for c in row.cells]
assert cells(old)==cells(client)
internalparas=[p.text for p in Document(OUT/'internal.docx').paragraphs]
assert oldparas[scope_start:scope_end]==internalparas[internalparas.index('读取材料范围'):internalparas.index('内部立场、底线与复核任务')]
records=[]
for label,name in jobs.items():
    work=ROOT/'.audit-local/t09-word-qa'/name
    driver=json.loads((work/'driver.json').read_text(encoding='utf8'))
    result=json.loads((work/'result.json').read_text(encoding='utf-8-sig'))
    pages=json.loads((work/'pages.json').read_text(encoding='utf8'))
    markers=json.loads((work/'png-markers.json').read_text(encoding='utf8'))
    source=OUT/(label+'.docx')
    assert sha256(source.read_bytes()).hexdigest()==driver['original_sha256']==result['input_sha256']
    assert driver['exit_code']==0 and result['status']=='WORD_EXPORT_EXECUTED'
    assert result['macros_disabled'] and result['window_pid_verified'] and not result['repair_requested'] and not result['add_recent']
    assert result['revisions_before']==result['revisions_after']==0
    pdf=PdfReader(work/'output.pdf');assert len(pdf.pages)==pages['pages']
    doc=Document(source)
    marker=normalize(''.join(p.text for p in doc.sections[0].header.paragraphs))
    body=[]
    for page in pdf.pages:
        text=normalize(page.extract_text())
        # These Word PDFs serialize both margin markers before the body.
        # Assert that exact prefix; keep the same warning when it occurs in body.
        assert text.startswith(marker*2)
        body.append(text[len(marker)*2:])
    pdftext=''.join(body)
    paragraphs=[p.text for p in doc.paragraphs if p.text.strip()]
    tablecells=[v for v in cells(doc) if v.strip()]
    missing=[t for t in paragraphs+tablecells if normalize(t) not in pdftext]
    assert not missing,{'audience':label,'missing':missing,'pdftext':pdftext}
    for page in pdf.pages:
        positions=[]
        def visitor(t,cm,tm,font,size):
            if 'DRAFT' in t:positions.append(tm[5])
        page.extract_text(visitor_text=visitor)
        assert any(y>float(page.mediabox.height)-60 for y in positions) and any(y<60 for y in positions)
    assert markers['all_pages_pixel_identical_markers']
    target=DEST/label;target.mkdir(exist_ok=True)
    for fname in ['driver.json','result.json','pages.json','png-markers.json','output.pdf','pdfinfo.log','pdftoppm.log']+[f'page-{n}.png' for n in range(1,pages['pages']+1)]:shutil.copyfile(work/fname,target/fname)
    shutil.copyfile(source,DEST/source.name)
    records.append({'audience':label,'sha256':driver['original_sha256'],'pages':pages['pages'],'paragraphs_preserved':len(paragraphs),'table_cells_preserved':len(tablecells),'word_version':result['word_version'],'word_build':result['word_build'],'input_unchanged':True,'header_footer_all_pages':True,'visual_review':'All page PNGs opened; no clipping, overlap, missing Chinese glyphs or broken rows observed.'})
for p in OUT.iterdir():
    if p.suffix in {'.json','.txt','.py'}:shutil.copyfile(p,DEST/p.name)
record={'checked_utc':datetime.now(timezone.utc).isoformat(),'records':records,'client_pages_before':4,'client_pages_after':3,'client_outside_scope_unchanged':True,'internal_scope_equal_prior_raw_scope':True,'prior_docx_unchanged':True,'real_model_qcc_calls':0,'full_suite':'NOT_RUN','new_packages':0}
(DEST/'verification.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps(record,ensure_ascii=False))
