from concurrent.futures import ThreadPoolExecutor
from threading import Event
from io import BytesIO
import pytest
from sqlalchemy import select, func
from cn_contract.domain import AssistedReviewInput
from cn_contract.review.storage import ReviewProposalRow
from cn_contract.security import ServiceError
from tests.integration.test_assisted_review import assisted_api
from tests.integration.test_case_bridge import selected
from tests.integration.test_plugin_candidates import input_body, open_write, post


def test_read_after_submit_does_not_destroy_replay(selected):
    client = selected[0]
    body = input_body(selected)
    _, auth = open_write(selected)
    first = post(client, auth, body)
    assert first.status_code == 201
    read = client.get('/api/case-bridge/context', headers=auth)
    replay = post(client, auth, body)
    assert replay.status_code == 201 and replay.json() == first.json(), (read.status_code, read.text, replay.status_code, replay.text)
    assert read.status_code == 200, read.text
    assert read.json()['context']['candidate_source'] == 'local_plugin_unverified'
    assert read.json()['context']['model_inference'] == 'NOT_ATTESTED'
    assert 'manual_synthetic' not in '\n'.join(read.json()['limitations'])


def test_write_grant_read_does_not_claim_manual_model_origin(selected):
    client = selected[0]
    _, auth = open_write(selected)
    response = client.get('/api/case-bridge/context', headers=auth)
    assert response.status_code == 200
    assert 'manual_synthetic' not in '\n'.join(response.json()['limitations'])
    assert '真实模型推理NOT_RUN' not in '\n'.join(response.json()['limitations'])


def test_commit_failure_leaves_no_ghost_receipt(selected, monkeypatch):
    client, database, _, bridge, _, matter, _, _ = selected
    body = input_body(selected)
    _, auth = open_write(selected)
    original_audit = bridge._audit
    original_commit = database.session_factory.class_.commit
    def audit(action, grant, actor, db):
        original_audit(action, grant, actor, db)
        if action == 'case_bridge_submitted': db.info['probe_fail_commit'] = True
    def commit(db):
        if db.info.pop('probe_fail_commit', False): raise ServiceError('STATE_CONFLICT', 409)
        return original_commit(db)
    monkeypatch.setattr(bridge, '_audit', audit)
    monkeypatch.setattr(database.session_factory.class_, 'commit', commit)
    assert post(client, auth, body).status_code == 409
    assert bridge.submission is None
    with database.session() as db:
        assert db.scalar(select(func.count()).select_from(ReviewProposalRow)) == 0
    monkeypatch.undo()
    assert post(client, auth, body).status_code == 201


def test_revocation_between_validation_and_write_blocks(selected, monkeypatch):
    from cn_contract.review import proposals
    client, database, _, bridge, _, _, _, _ = selected
    value = AssistedReviewInput.model_validate(input_body(selected))
    _, auth = open_write(selected)
    arrived, release = Event(), Event()
    original = proposals._context_state
    def pause(*args, **kwargs):
        result = original(*args, **kwargs)
        arrived.set()
        assert release.wait(5)
        return result
    monkeypatch.setattr(proposals, '_context_state', pause)
    def submit():
        with database.session() as db:
            return bridge.submit(auth['Authorization'], value, db)
    with ThreadPoolExecutor(max_workers=1) as pool:
        pending = pool.submit(submit)
        assert arrived.wait(5)
        bridge.close()
        release.set()
        with pytest.raises(ServiceError) as error: pending.result(timeout=5)
        assert error.value.status_code == 401
    with database.session() as db:
        assert db.scalar(select(func.count()).select_from(ReviewProposalRow)) == 0


def test_reopened_database_retains_receipt_but_not_grant(selected):
    from cn_contract.persistence import Database
    from cn_contract.web.case_bridge import CaseBridge
    from cn_contract.web.app import create_app, initialize_storage
    from fastapi.testclient import TestClient
    client, database, _, bridge, path, matter, _, _ = selected
    body = input_body(selected)
    _, auth = open_write(selected)
    first = post(client, auth, body)
    assert first.status_code == 201
    reopened = Database(str(database.engine.url), database.storage_root, database.settings, initialize_schema=False)
    initialize_storage(reopened)
    app = create_app(database.settings, database=reopened, purpose='assisted_draft', assisted_policy=database._assisted_policy)
    from cn_contract.web.case_bridge import attach_case_bridge
    attach_case_bridge(app, reopened, path)
    with TestClient(app) as new:
        new.cookies.update(client.cookies)
        assert post(new, auth, body).status_code == 401
        run = new.get(f'/api/matters/{matter}/reviews/{first.json()["run_id"]}')
        assert run.status_code == 200, run.text
        assert run.json()['proposal']['execution_receipt']['source'] == 'local_plugin_unverified'
    reopened.engine.dispose()


def test_lawyer_disposition_preserves_replay_and_native_redline_is_not_manual_claim(selected):
    from docx import Document
    client, _, _, _, _, matter, _, csrf = selected
    body = input_body(selected)
    _, auth = open_write(selected)
    receipt = post(client, auth, body)
    assert receipt.status_code == 201
    run = client.get(f'/api/matters/{matter}/reviews/{receipt.json()["run_id"]}').json()
    for finding in run['findings']:
        response = client.post(f'/api/matters/{matter}/findings/{finding["id"]}/decision', headers=csrf,
            json={'run_id':run['id'], 'run_sha256':run['run_sha256'], 'decision':'accept', 'reason':'独立虚构复核'})
        assert response.status_code == 201, response.text
    again = post(client, auth, body)
    assert again.status_code == 201 and again.json() == receipt.json()
    finding = run['findings'][0]
    patch = client.post(f'/api/matters/{matter}/patches/approve', headers={**csrf, 'Idempotency-Key':'review-probe-patch'},
        json={'run_id':run['id'], 'run_sha256':run['run_sha256'], 'finding_id':finding['id'],
            'anchor_id':finding['anchor_ids'][0], 'new_text':'独立虚构律师改写，用于工程探针。', 'reason':'测试律师明确批准', 'audiences':['internal']})
    assert patch.status_code == 201, patch.text
    preview = client.post(f'/api/matters/{matter}/exports/preview', headers=csrf,
        json={'run_id':run['id'], 'run_sha256':run['run_sha256'], 'audience':'internal', 'patch_ids':[patch.json()['id']]})
    assert preview.status_code == 201, preview.text
    native = next(a for a in preview.json()['artifacts'] if a['type'] == 'redline_contract_docx')
    response = client.get(f'/api/matters/{matter}/exports/{native["id"]}/download')
    assert response.status_code == 200, response.text
    doc = Document(BytesIO(response.content))
    wording = '\n'.join(s.header.paragraphs[0].text + s.footer.paragraphs[0].text for s in doc.sections)
    assert 'DRAFT MOCK' in wording and '不构成正式法律意见' in wording
    assert not any(s in wording for s in ('manual_synthetic', '人工/合成', 'NOT_RUN', '推理未运行'))


def test_read_during_submit_does_not_use_precommit_finished_flag(selected, monkeypatch):
    from threading import current_thread
    from cn_contract.review import proposals
    client, database, _, bridge, _, _, _, _ = selected
    body = input_body(selected)
    _, auth = open_write(selected)
    writing, release_write, read_context = Event(), Event(), Event()
    original_audit, original_context = bridge._audit, proposals._context_state
    def audit(action, *args):
        original_audit(action, *args)
        if action == 'case_bridge_submitted':
            writing.set()
            assert release_write.wait(5)
    def context(*args, **kwargs):
        if current_thread().name.startswith('reader'):
            read_context.set()
        return original_context(*args, **kwargs)
    monkeypatch.setattr(bridge, '_audit', audit)
    monkeypatch.setattr(proposals, '_context_state', context)
    def read():
        with database.session() as db:
            return bridge.read(auth['Authorization'], db)
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix='writer') as writer, ThreadPoolExecutor(max_workers=1, thread_name_prefix='reader') as reader:
        submitted = writer.submit(post, client, auth, body)
        assert writing.wait(5)
        pending_read = reader.submit(read)
        read_context.wait(1)
        release_write.set()
        first = submitted.result(timeout=5)
        assert first.status_code == 201, first.text
        result = pending_read.result(timeout=5)
        assert result
    replay = post(client, auth, body)
    assert replay.status_code == 201 and replay.json() == first.json()


def test_identical_replay_waits_for_postcommit_receipt_publication(selected, monkeypatch):
    from threading import current_thread
    from cn_contract.review import proposals
    client, database, _, bridge, _, _, _, _ = selected
    value = AssistedReviewInput.model_validate(input_body(selected))
    _, auth = open_write(selected)
    committed, publish, second_context = Event(), Event(), Event()
    original_commit = database.session_factory.class_.commit
    original_context = proposals._context_state
    def commit(db):
        result = original_commit(db)
        if current_thread().name.startswith('first') and not committed.is_set():
            committed.set()
            assert publish.wait(5)
        return result
    def context(*args, **kwargs):
        if current_thread().name.startswith('second'):
            second_context.set()
        return original_context(*args, **kwargs)
    monkeypatch.setattr(database.session_factory.class_, 'commit', commit)
    monkeypatch.setattr(proposals, '_context_state', context)
    def submit():
        with database.session() as db:
            return bridge.submit(auth['Authorization'], value, db)
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix='first') as first_pool, ThreadPoolExecutor(max_workers=1, thread_name_prefix='second') as second_pool:
        first = first_pool.submit(submit)
        assert committed.wait(5)
        second = second_pool.submit(submit)
        second_context.wait(1)
        publish.set()
        assert first.result(timeout=5) == second.result(timeout=5)
