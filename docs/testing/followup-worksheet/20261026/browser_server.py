from pathlib import Path
from tempfile import mkdtemp
from io import BytesIO
from uuid import uuid4
import hashlib,json
from docx import Document
import uvicorn
from cn_contract import domain as d
from cn_contract.web.local_server import initialize_workspace,build_local_app
from cn_contract.assistant import service,profile_store
from cn_contract.security import login,authenticate
base=Path(mkdtemp(prefix='qx-followup-browser-',dir='/private/tmp'));root=base/'workspaces/main'
initialize_workspace(root,'合成跟进律师','fiction123',data_parent=base);app=build_local_app(root,port=18772,data_parent=base)
doc=Document();doc.add_paragraph('独立虚构交付记录：甲向乙转账100000元。');doc.add_paragraph('乙的相反说法：其中部分是其他往来款。');s=BytesIO();doc.save(s);raw=s.getvalue();(base/'合成交付.docx').write_bytes(raw)
app.state.database._assistant_synthetic_hashes=frozenset({hashlib.sha256(raw).hexdigest()})
with app.state.database.session() as db:
 token,_=login('合成跟进律师','fiction123',db);actor=authenticate(token,db)
 item=service.create_conversation(actor,d.AssistantConversationCreateInput(request_id=uuid4(),display_title='合成借贷补查跟进'),db)
profile_store.save_profile(actor,d.ProfileCommitInput(request_id=uuid4(),patch=d.ProfilePatch(scope=item.scope,base_profile_revision=0,changes=()),manual_edits=(d.ProfileEdit(field='consultation',value={'kind':'text','value':'甲称借出10万元，乙称部分已还，性质和还款金额待核。'}),d.ProfileEdit(field='internal_notes',value={'kind':'text','value':'合成内部策略，不对客户展示。'}))),database=app.state.database)
(base/'fixture.json').write_text(json.dumps({'cid':str(item.scope.conversation_id),'root':str(root)},ensure_ascii=False));Path('/private/tmp/qx-followup-browser-path').write_text(str(base))
uvicorn.run(app,host='127.0.0.1',port=18772,log_level='warning')
