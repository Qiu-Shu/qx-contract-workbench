"""Explicit candidate-write grant; source provenance and lawyer control remain."""
from copy import deepcopy
import json
from uuid import UUID, uuid4
from datetime import datetime, timedelta, timezone

import pytest
from sqlalchemy import select, func
from cn_contract.review.storage import ReviewProposalRow
from tests.integration.test_case_bridge import selected
from tests.integration.test_assisted_review import assisted_api, proposal_body, CASES


def input_body(selected):
    client, _, _, _, _, matter, context, _ = selected
    identifier, digest = next(iter(context['document_hashes'].items()))
    parsed = client.get(f'/api/matters/{matter}/documents/{identifier}/anchors').json()
    facts = client.get(f'/api/matters/{matter}/facts').json()
    return proposal_body(CASES[0], {'version_id': identifier, 'sha256': digest}, parsed, facts, context)


def open_write(selected, allow=True):
    client, _, _, _, path, matter, context, csrf = selected
    body = {'context_sha256': context['context_sha256'], 'share_with_chat': True}
    if allow: body['allow_candidate_submission'] = True
    response = client.post(f'/api/matters/{matter}/case-bridge', headers=csrf, json=body)
    assert response.status_code == 201, response.text
    token = json.loads(path.read_text('utf8'))['token']
    return response.json(), {'Authorization': 'Bearer ' + token}


def post(client, auth, body):
    return client.post('/api/case-bridge/candidates', headers=auth, json=body)


def test_read_only_grant_cannot_write(selected):
    client = selected[0]
    body = input_body(selected)
    _, auth = open_write(selected, False)
    response = post(client, auth, body)
    assert response.status_code == 403, response.text
    assert client.get('/api/case-bridge/context', headers=auth).status_code == 200


def test_plugin_source_persists_as_unverified_and_awaits_lawyer(selected):
    client, database, _, _, _, matter, _, _ = selected
    body = input_body(selected)
    grant, auth = open_write(selected)
    assert grant['read_only'] is False
    response = post(client, auth, body)
    assert response.status_code == 201, response.text
    receipt = response.json()
    assert receipt['grant_id'] == grant['id'] and receipt['source'] == 'local_plugin_unverified'
    run = client.get(f'/api/matters/{matter}/reviews/{receipt["run_id"]}')
    assert run.status_code == 200, run.text
    assert all(f['human_decision'] is None for f in run.json()['findings'])
    assert all('插件提交' in f['business_effect'] for f in run.json()['findings'])
    assert all('离线人工/合成' not in f['business_effect'] for f in run.json()['findings'])
    assert any('未核验' in line for line in run.json()['limitations'])
    assert client.get(f'/api/matters/{matter}').json()['matter']['status'] == 'awaiting_lawyer'
    with database.session() as db:
        rows = list(db.scalars(select(ReviewProposalRow)))
        assert len(rows) == 1
        stored = json.loads(rows[0].payload)['execution_receipt']
        assert stored['source'] == 'local_plugin_unverified' and stored['model_identity'] == 'UNVERIFIED'
        assert stored['grant_id'] == grant['id']


def test_one_grant_one_round_and_same_input_replays_receipt(selected):
    client, database, _, _, _, _, _, _ = selected
    body = input_body(selected); _, auth = open_write(selected)
    first = post(client, auth, body)
    assert first.status_code == 201, first.text
    again = post(client, auth, body)
    assert again.status_code == 201 and again.json() == first.json()
    changed = deepcopy(body); changed['opinions'][0]['analysis'] += ' changed'
    assert post(client, auth, changed).status_code == 409
    with database.session() as db: assert db.scalar(select(func.count()).select_from(ReviewProposalRow)) == 1


@pytest.mark.parametrize('change', ['expired', 'revoked', 'version', 'wrong_anchor', 'fake_receipt', 'fake_actor', 'query'])
def test_stale_or_untrusted_submission_is_rejected(selected, change):
    from cn_contract.security import remove_membership
    client, database, ids, bridge, _, matter, _, csrf = selected
    body = input_body(selected); _, auth = open_write(selected)
    if change == 'expired': bridge.now = lambda: datetime.now(timezone.utc) + timedelta(minutes=16)
    elif change == 'revoked':
        with database.session() as db: remove_membership(UUID(matter), ids['lawyer'], db)
    elif change == 'version': body['context_sha256'] = '0'*64
    elif change == 'wrong_anchor': body['opinions'][0]['anchor_ids'] = [str(uuid4())]
    elif change == 'fake_receipt': body['source'] = 'verified_chatgpt'
    elif change == 'fake_actor': body['actor_id'] = str(ids['reviewer'])
    if change == 'query': response = client.post('/api/case-bridge/candidates?matter_id='+matter, headers=auth, json=body)
    else: response = post(client, auth, body)
    assert response.status_code in (401, 403, 409, 422), response.text
    with database.session() as db: assert db.scalar(select(func.count()).select_from(ReviewProposalRow)) == 0


def test_manual_endpoint_cannot_claim_plugin_provenance(selected):
    client, _, _, _, _, matter, _, csrf = selected
    body = input_body(selected);body['execution_receipt'] = {'source': 'local_plugin_unverified'}
    response = client.post(f'/api/matters/{matter}/review-proposals', headers={**csrf,'Idempotency-Key':'fake-source'}, json=body)
    assert response.status_code == 422


def test_write_consent_is_strict_boolean(selected):
    client, _, _, _, _, matter, context, csrf = selected
    response = client.post(f'/api/matters/{matter}/case-bridge', headers=csrf,
        json={'context_sha256': context['context_sha256'], 'share_with_chat': True, 'allow_candidate_submission': 'true'})
    assert response.status_code == 422


def test_concurrent_identical_submission_persists_once(selected):
    from concurrent.futures import ThreadPoolExecutor
    client, database, _, bridge, _, _, _, _ = selected
    from cn_contract.domain import AssistedReviewInput
    body = AssistedReviewInput.model_validate(input_body(selected)); _, auth = open_write(selected)
    def submit():
        with database.session() as db: return bridge.submit(auth['Authorization'], body, db)
    with ThreadPoolExecutor(max_workers=2) as pool:
        first, second = list(pool.map(lambda _: submit(), range(2)))
    assert first == second
    with database.session() as db: assert db.scalar(select(func.count()).select_from(ReviewProposalRow)) == 1


def test_failed_transaction_leaves_no_receipt_or_candidate(selected, monkeypatch):
    from cn_contract.security import ServiceError
    client, database, _, bridge, _, matter, _, _ = selected
    body=input_body(selected); _, auth=open_write(selected)
    original=bridge._audit
    def fail(action,*args):
        if action=='case_bridge_submitted': raise ServiceError('STATE_CONFLICT',409)
        return original(action,*args)
    monkeypatch.setattr(bridge,'_audit',fail)
    assert post(client,auth,body).status_code==409
    assert bridge.submission is None
    with database.session() as db: assert db.scalar(select(func.count()).select_from(ReviewProposalRow))==0
    assert client.get(f'/api/matters/{matter}').json()['matter']['status']=='ready_for_review'
    monkeypatch.setattr(bridge,'_audit',original)
    assert post(client,auth,body).status_code==201


def test_replay_after_revocation_is_denied(selected):
    client, _, _, _, _, matter, _, csrf=selected
    body=input_body(selected); _,auth=open_write(selected)
    assert post(client,auth,body).status_code==201
    assert client.delete(f'/api/matters/{matter}/case-bridge',headers=csrf).status_code==204
    assert post(client,auth,body).status_code==401


def test_lawyer_disposition_and_export_preserve_plugin_origin(selected):
    from io import BytesIO
    from docx import Document
    from tests.integration.test_assisted_review import docx_text
    client, _, _, _, _, matter, _, csrf=selected
    body=input_body(selected); _,auth=open_write(selected)
    submitted=post(client,auth,body)
    assert submitted.status_code==201,submitted.text
    run=client.get(f'/api/matters/{matter}/reviews/{submitted.json()["run_id"]}').json()
    for f in run['findings']:
        response=client.post(f'/api/matters/{matter}/findings/{f["id"]}/decision',headers=csrf,
            json={'run_id':run['id'],'run_sha256':run['run_sha256'],'decision':'accept','reason':'独立虚构来源与候选已由测试律师核对'})
        assert response.status_code==201,response.text
    preview=client.post(f'/api/matters/{matter}/exports/preview',headers=csrf,
        json={'run_id':run['id'],'run_sha256':run['run_sha256'],'audience':'client','patch_ids':[]})
    assert preview.status_code==201,preview.text
    identifier=preview.json()['artifacts'][0]['id']
    output=client.get(f'/api/matters/{matter}/exports/{identifier}/download')
    assert output.status_code==200
    doc=Document(BytesIO(output.content))
    assert all('插件候选' in s.header.paragraphs[0].text and '未核验' in s.footer.paragraphs[0].text for s in doc.sections)
    assert doc.core_properties.keywords=='DRAFT PLUGIN-UNVERIFIED'
    assert '离线人工/合成候选' not in docx_text(output.content)


def test_actual_stdio_candidate_submission_over_loopback(selected):
    import socket,time
    from threading import Thread
    import uvicorn
    from tests.helpers.offline_mcp_child import child,initialize
    client,database,_,_,path,matter,_,_=selected
    body=input_body(selected);open_write(selected)
    sock=socket.socket();sock.bind(('127.0.0.1',8767))
    server=uvicorn.Server(uvicorn.Config(client.app,host='127.0.0.1',port=8767,lifespan='off',access_log=False,log_level='error'))
    thread=Thread(target=lambda:server.run(sockets=[sock]),daemon=True);thread.start()
    try:
        deadline=time.monotonic()+10
        while not server.started and time.monotonic()<deadline:time.sleep(.02)
        assert server.started
        code=('from pathlib import Path; from cn_contract.mcp import case_client,offline_server; '
              f'case_client.CREDENTIAL_PATH=Path({str(path)!r});offline_server.main()')
        with child(launch_code=code) as (process,send,stderr):
            initialize(send)
            result=send('tools/call',{'name':'submit_selected_candidates','arguments':{'candidate':body}},request_id=2)
            assert 'result' in result,result
            receipt=result['result']['structuredContent']
            assert receipt['source']=='local_plugin_unverified' and receipt['matter_id']==matter
            replay=send('tools/call',{'name':'submit_selected_candidates','arguments':{'candidate':body}},request_id=3)
            assert replay['result']['structuredContent']==receipt
            bad=send('tools/call',{'name':'submit_selected_candidates','arguments':{'candidate':body,'actor':'fake'}},request_id=4)
            assert bad['error']['message']=='OFFLINE_REQUEST_REJECTED'
        assert process.returncode==0 and stderr==[]
        with database.session() as db: assert db.scalar(select(func.count()).select_from(ReviewProposalRow))==1
    finally:
        server.should_exit=True;thread.join(10);sock.close();assert not thread.is_alive()


def test_read_after_submit_does_not_destroy_replay(selected):
    client = selected[0]
    body = input_body(selected)
    _, auth = open_write(selected)
    first = post(client, auth, body)
    assert first.status_code == 201
    read = client.get('/api/case-bridge/context', headers=auth)
    replay = post(client, auth, body)
    assert replay.status_code == 201 and replay.json() == first.json(), (read.status_code, read.text, replay.status_code, replay.text)
    assert read.status_code == 200, read.text
    assert read.json()['context']['candidate_source'] == 'local_plugin_unverified'
    assert read.json()['context']['model_inference'] == 'NOT_ATTESTED'
    assert 'manual_synthetic' not in '\n'.join(read.json()['limitations'])


def test_write_grant_read_does_not_claim_manual_model_origin(selected):
    client = selected[0]
    _, auth = open_write(selected)
    response = client.get('/api/case-bridge/context', headers=auth)
    assert response.status_code == 200
    assert 'manual_synthetic' not in '\n'.join(response.json()['limitations'])
    assert '真实模型推理NOT_RUN' not in '\n'.join(response.json()['limitations'])




def test_read_during_submit_does_not_use_precommit_finished_flag(selected, monkeypatch):
    from threading import current_thread, Event
    from concurrent.futures import ThreadPoolExecutor
    from cn_contract.review import proposals
    client, database, _, bridge, _, _, _, _ = selected
    body = input_body(selected)
    _, auth = open_write(selected)
    writing, release_write, read_context = Event(), Event(), Event()
    original_audit, original_context = bridge._audit, proposals._context_state
    def audit(action, *args):
        original_audit(action, *args)
        if action == 'case_bridge_submitted':
            writing.set()
            assert release_write.wait(5)
    def context(*args, **kwargs):
        if current_thread().name.startswith('reader'):
            read_context.set()
        return original_context(*args, **kwargs)
    monkeypatch.setattr(bridge, '_audit', audit)
    monkeypatch.setattr(proposals, '_context_state', context)
    def read():
        with database.session() as db:
            return bridge.read(auth['Authorization'], db)
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix='writer') as writer, ThreadPoolExecutor(max_workers=1, thread_name_prefix='reader') as reader:
        submitted = writer.submit(post, client, auth, body)
        assert writing.wait(5)
        pending_read = reader.submit(read)
        read_context.wait(1)
        release_write.set()
        first = submitted.result(timeout=5)
        assert first.status_code == 201, first.text
        result = pending_read.result(timeout=5)
        assert result
    replay = post(client, auth, body)
    assert replay.status_code == 201 and replay.json() == first.json()


def test_identical_replay_waits_for_postcommit_receipt_publication(selected, monkeypatch):
    from threading import current_thread, Event
    from concurrent.futures import ThreadPoolExecutor
    from cn_contract.domain import AssistedReviewInput
    from cn_contract.review import proposals
    client, database, _, bridge, _, _, _, _ = selected
    value = AssistedReviewInput.model_validate(input_body(selected))
    _, auth = open_write(selected)
    committed, publish, second_context = Event(), Event(), Event()
    original_commit = database.session_factory.class_.commit
    original_context = proposals._context_state
    def commit(db):
        result = original_commit(db)
        if current_thread().name.startswith('first') and not committed.is_set():
            committed.set()
            assert publish.wait(5)
        return result
    def context(*args, **kwargs):
        if current_thread().name.startswith('second'):
            second_context.set()
        return original_context(*args, **kwargs)
    monkeypatch.setattr(database.session_factory.class_, 'commit', commit)
    monkeypatch.setattr(proposals, '_context_state', context)
    def submit():
        with database.session() as db:
            return bridge.submit(auth['Authorization'], value, db)
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix='first') as first_pool, ThreadPoolExecutor(max_workers=1, thread_name_prefix='second') as second_pool:
        first = first_pool.submit(submit)
        assert committed.wait(5)
        second = second_pool.submit(submit)
        second_context.wait(1)
        publish.set()
        assert first.result(timeout=5) == second.result(timeout=5)
