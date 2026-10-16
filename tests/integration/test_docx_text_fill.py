"""Independent python-docx fixtures for bounded Word 2010 text effects."""
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
import pytest

from cn_contract.facts.storage import source_anchor_is_eligible
from tests.integration.test_docx_office_parts import W, edited, parse, saved

W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'
FILL = ('<w14:textFill xmlns:w14="' + W14 + '"><w14:solidFill>'
        '<w14:schemeClr w14:val="accent1"><w14:lumMod w14:val="50000"/>'
        '<w14:lumOff w14:val="50000"/></w14:schemeClr></w14:solidFill></w14:textFill>')


def filled_document(location='paragraph', fill=FILL, part='word/styles.xml'):
    document = Document()
    paragraph = document.add_paragraph('SYNTHETIC fill candidate 18500 CNY.')
    document.add_paragraph('SYNTHETIC unaffected paragraph.')
    style_id = None
    if location in {'paragraph', 'character', 'unused', 'table', 'conditional_table'}:
        kind = (WD_STYLE_TYPE.CHARACTER if location == 'character' else
                WD_STYLE_TYPE.TABLE if 'table' in location else WD_STYLE_TYPE.PARAGRAPH)
        base = document.styles.add_style('SyntheticFillBase', kind)
        derived = document.styles.add_style('SyntheticFillDerived', kind)
        derived.base_style = base
        style_id = base.style_id
        if location == 'paragraph':
            paragraph.style = derived
        elif location == 'character':
            paragraph.runs[0].style = derived
        elif 'table' in location:
            paragraph.text = 'SYNTHETIC ordinary body.'
            table = document.add_table(rows=1, cols=1)
            table.style = derived
            table.cell(0, 0).text = 'SYNTHETIC fill candidate 18500 CNY.'
    elif location == 'default_paragraph':
        style_id = 'Normal'
    elif location == 'default_character':
        style_id = 'DefaultParagraphFont'
    elif location == 'default_table':
        style_id = 'TableNormal'
        paragraph.text = 'SYNTHETIC ordinary body.'
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'SYNTHETIC fill candidate 18500 CNY.'
    elif location == 'direct':
        paragraph.runs[0]._r.get_or_add_rPr()
        part = 'word/document.xml'

    def add_fill(root):
        if location == 'direct':
            props = root.find(f'{{{W}}}body/{{{W}}}p/{{{W}}}r/{{{W}}}rPr')
        elif location == 'doc_defaults':
            props = root.find(f'{{{W}}}docDefaults/{{{W}}}rPrDefault/{{{W}}}rPr')
        else:
            style = next(node for node in root if node.get(f'{{{W}}}styleId') == style_id)
            if location == 'conditional_table':
                style = ET.SubElement(style, f'{{{W}}}tblStylePr', {f'{{{W}}}type': 'firstRow'})
            props = style.find(f'{{{W}}}rPr')
            if props is None:
                props = ET.SubElement(style, f'{{{W}}}rPr')
        props.append(ET.fromstring(fill))

    data = saved(document)
    if part == 'word/stylesWithEffects.xml' and style_id and style_id.startswith('Synthetic'):
        # python-docx's effects part is a template; copy only our synthetic styles.
        from io import BytesIO
        from zipfile import ZipFile
        with ZipFile(BytesIO(data)) as package:
            styles = ET.fromstring(package.read('word/styles.xml'))
        def copy_styles(root):
            root.extend(node for node in styles if (node.get(f'{{{W}}}styleId') or '').startswith('Synthetic'))
        data = edited(data, {part: copy_styles})
    return edited(data, {part: add_fill})


@pytest.mark.parametrize('part', ['word/styles.xml', 'word/stylesWithEffects.xml'])
@pytest.mark.parametrize('location', [
    'doc_defaults', 'default_paragraph', 'default_character', 'paragraph', 'character',
    'table', 'default_table', 'conditional_table',
])
def test_used_text_fill_is_local_visibility_gap_not_untrusted_styles(location, part):
    result = parse(filled_document(location, part=part))
    candidate = next(anchor for anchor in result.anchors if 'fill candidate' in anchor.quote)
    assert not source_anchor_is_eligible(candidate, result)
    assert any('DOCX_VISIBILITY_UNCERTAIN:' in issue for issue in result.quality_issues)
    assert not any('style_dependency' in gap for gap in result.unsupported_parts)
    assert f'{part}:metadata_read' in result.parts
    if location not in {'doc_defaults', 'default_paragraph', 'default_character'}:
        ordinary = next(anchor for anchor in result.anchors if 'unaffected paragraph' in anchor.quote)
        assert source_anchor_is_eligible(ordinary, result)


@pytest.mark.parametrize('part', ['word/styles.xml', 'word/stylesWithEffects.xml'])
def test_unused_legal_text_fill_does_not_block_plain_body(part):
    result = parse(filled_document('unused', part=part))
    assert len(result.anchors) == 2
    assert all(source_anchor_is_eligible(anchor, result) for anchor in result.anchors)
    assert result.unsupported_parts == result.quality_issues == ()


@pytest.mark.parametrize('children', [
    '', '<w14:noFill/>', '<w14:solidFill/>',
    '<w14:solidFill><w14:srgbClr w14:val="000000"/></w14:solidFill>',
    '<w14:solidFill><w14:srgbClr w14:val="FFFFFF"/></w14:solidFill>',
    '<w14:solidFill><w14:srgbClr w14:val="000000"><w14:alpha w14:val="0"/></w14:srgbClr></w14:solidFill>',
    '<w14:solidFill><w14:schemeClr w14:val="tx1"/></w14:solidFill>',
])
@pytest.mark.parametrize('location', ['paragraph', 'direct'])
def test_known_fill_variants_remain_uncertain_until_rendered(children, location):
    fill = f'<w14:textFill xmlns:w14="{W14}">{children}</w14:textFill>'
    result = parse(filled_document(location, fill))
    assert len(result.anchors) == 2
    assert not source_anchor_is_eligible(result.anchors[0], result)
    assert source_anchor_is_eligible(result.anchors[1], result)
    assert any(':visibility:' in gap for gap in result.unsupported_parts)
    assert not any('style_dependency' in gap for gap in result.unsupported_parts)


@pytest.mark.parametrize('mutation', [
    'unknown_effect', 'unknown_transform', 'wrong_namespace', 'wrong_color_namespace',
    'invalid_scheme', 'invalid_rgb', 'missing_value', 'invalid_alpha', 'invalid_number',
    'nested_text', 'two_colors', 'unknown_attribute', 'unqualified_value', 'misplaced_fill',
])
def test_unknown_or_malformed_style_effect_never_yields_eligible_text(mutation):
    fill = ET.fromstring(FILL)
    color = fill[0][0]
    if mutation == 'unknown_effect':
        fill.tag = f'{{{W14}}}unrecognizedEffect'
    elif mutation == 'unknown_transform':
        color[0].tag = f'{{{W14}}}unrecognizedTransform'
    elif mutation == 'wrong_namespace':
        fill.tag = '{urn:synthetic:unrecognized}textFill'
    elif mutation == 'wrong_color_namespace':
        color.tag = f'{{{W}}}schemeClr'
    elif mutation == 'invalid_scheme':
        color.set(f'{{{W14}}}val', 'unrecognizedTheme')
    elif mutation == 'invalid_rgb':
        color.tag = f'{{{W14}}}srgbClr'
        color.set(f'{{{W14}}}val', 'ZZZZZZ')
    elif mutation == 'missing_value':
        color.attrib.clear()
    elif mutation == 'invalid_alpha':
        color[0].tag = f'{{{W14}}}alpha'
        color[0].set(f'{{{W14}}}val', '100001')
    elif mutation == 'invalid_number':
        color[0].set(f'{{{W14}}}val', 'not-a-percentage')
    elif mutation == 'nested_text':
        color[0].text = 'SYNTHETIC unexplained obligation'
    elif mutation == 'two_colors':
        ET.SubElement(fill[0], f'{{{W14}}}srgbClr', {f'{{{W14}}}val': '000000'})
    elif mutation == 'unknown_attribute':
        fill.set('unrecognized', '1')
    elif mutation == 'unqualified_value':
        color.attrib.clear()
        color.set('val', 'accent1')
    elif mutation == 'misplaced_fill':
        wrapper = ET.Element(f'{{{W}}}pPr')
        wrapper.append(fill)
        fill = wrapper
    result = parse(filled_document('paragraph', ET.tostring(fill)))
    assert not any(source_anchor_is_eligible(anchor, result) for anchor in result.anchors)
    assert 'word/styles.xml:unsupported' in result.parts


@pytest.mark.parametrize('effect', ['textFill', 'unrecognizedEffect'])
def test_unparsed_direct_effect_is_never_silently_skipped(effect):
    fill = f'<w14:{effect} xmlns:w14="{W14}"><w14:unrecognizedChild/></w14:{effect}>'
    result = parse(filled_document('direct', fill))
    assert len(result.anchors) == 2
    assert not source_anchor_is_eligible(result.anchors[0], result)
    assert source_anchor_is_eligible(result.anchors[1], result)


def test_wrong_namespace_direct_fill_is_not_treated_as_ordinary_run_formatting():
    fill = (f'<w:textFill xmlns:w="{W}"><w:noFill/></w:textFill>')
    result = parse(filled_document('direct', fill))
    assert len(result.anchors) == 2
    assert not source_anchor_is_eligible(result.anchors[0], result)
    assert source_anchor_is_eligible(result.anchors[1], result)


@pytest.mark.parametrize('mutation', ['duplicate', 'wrong_parent', 'multiple_fills', 'gradient'])
def test_ambiguous_or_unsupported_fill_structure_keeps_styles_manual(mutation):
    data = filled_document()
    def change(root):
        style = next(node for node in root if node.get(f'{{{W}}}styleId') == 'SyntheticFillBase')
        props = style.find(f'{{{W}}}rPr')
        fill = props.find(f'{{{W14}}}textFill')
        if mutation == 'duplicate':
            props.append(ET.fromstring(FILL))
        elif mutation == 'wrong_parent':
            props.remove(fill)
            style.append(fill)
        elif mutation == 'multiple_fills':
            ET.SubElement(fill, f'{{{W14}}}noFill')
        else:
            fill[0].tag = f'{{{W14}}}gradFill'
    result = parse(edited(data, {'word/styles.xml': change}))
    assert not any(source_anchor_is_eligible(anchor, result) for anchor in result.anchors)
    assert 'word/styles.xml:unsupported' in result.parts


def test_supported_fill_does_not_override_hidden_style():
    data = filled_document()
    def hide(root):
        style = next(node for node in root if node.get(f'{{{W}}}styleId') == 'SyntheticFillBase')
        ET.SubElement(style.find(f'{{{W}}}rPr'), f'{{{W}}}vanish')
    result = parse(edited(data, {'word/styles.xml': hide}))
    assert [anchor.quote for anchor in result.anchors] == ['SYNTHETIC unaffected paragraph.']
    assert any(':hidden:' in gap for gap in result.unsupported_parts)
