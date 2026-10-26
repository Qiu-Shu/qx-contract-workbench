"""Source-bound evidence rows on independent fictional native workspaces."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
import pytest
from sqlalchemy import text

from test_assistant_routes import native, chat, create, sign_in
from test_reception_worksheet import save


def begin(client,h):
    item=create(client,h);base='/api/assistant/conversations/'+item['scope']['conversation_id']
    return base,base+'/evidence'


def request(view,rows=()):
    return {'request_id':str(uuid4()),'expected_scope':view['scope'],
        'expected_profile_revision':view['profile_revision'],'base_revision':view['revision'],'rows':list(rows)}


def statement(client,h,base,url):
    save(client,h,base+'/reception-worksheet',consultation='甲称借出10万元，乙称已经部分还款。')
    v=client.get(url).json();quote=v['statement']
    return {'id':str(uuid4()),'category':'agreement','content':'甲自述借出10万元，尚待材料核对',
        'event_date':None,'review_state':'pending','note':'不得直接推定余额',
        'source':{'kind':'statement','profile_revision':v['profile_revision'],'start':0,'end':len(quote),'quote':quote}}


def material(client,h,base,database):
    doc=Document();doc.add_paragraph('独立虚构交付记录：甲向乙转账100000元。');doc.add_paragraph('乙的相反说法：其中部分是其他往来款。')
    stream=BytesIO();doc.save(stream);raw=stream.getvalue()
    database._assistant_synthetic_hashes=frozenset({sha256(raw).hexdigest()})
    v=client.get(base).json();body={'request_id':str(uuid4()),'expected_scope':v['scope'],
        'original_name':'合成交付.docx','mime':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'policy':{'local_processing_accepted':True}}
    r=client.post(base+'/attachments',headers=h,files=[('file',(body['original_name'],raw,body['mime'])),('metadata',(None,json.dumps(body)))])
    assert r.status_code==201,r.text
    attachment=r.json()['attachment'];v=client.get(base).json()
    r=client.post(base+'/attachments/'+attachment['id']+'/parses',headers=h,json={'request_id':str(uuid4()),'expected_scope':v['scope']})
    assert r.status_code==200,r.text
    parsed=r.json()['parse'];anchor=parsed['result']['anchors'][0]
    return attachment,parsed,{'kind':'document','conversation_id':v['scope']['conversation_id'],'attachment_id':attachment['id'],
        'original_sha256':attachment['sha256'],'parse_revision':parsed['revision'],'result_sha256':parsed['result_sha256'],'anchor':anchor}


def test_statement_edit_history_retry_and_no_model(chat):
    client,database,_=chat;h=sign_in(client);base,url=begin(client,h)
    r=client.get(url);assert r.status_code==200,r.text
    assert r.json()['revision']==0 and r.json()['rows']==[]
    row=statement(client,h,base,url);body=request(client.get(url).json(),[row])
    r=client.post(url+'/commits',headers=h,json=body);assert r.status_code==200,r.text
    first=r.json();assert first['revision']==1
    edited=deepcopy(row);edited['content']='只记录客户陈述，交付待核';edited['review_state']='disputed'
    second=client.post(url+'/commits',headers=h,json=request(client.get(url).json(),[edited]));assert second.status_code==200,second.text
    assert client.get(url).json()['rows']==[edited]
    assert client.get(url+'/revisions/1').json()['rows']==[row]
    assert client.post(url+'/commits',headers=h,json=body).json()==first
    assert client.get(url+'/receipts/'+body['request_id']).json()==first
    assert client.get(url+'/receipts/'+str(uuid4())).json() is None
    bad=deepcopy(body);bad['rows'][0]['note']='同号不同内容';assert client.post(url+'/commits',headers=h,json=bad).status_code==409
    with database.session() as db:
        assert db.scalar(text('SELECT count(*) FROM assistant_model_attempts'))==0
        assert db.scalar(text('SELECT count(*) FROM matters'))==0


def test_document_source_original_and_manual_excerpt_are_distinct(chat):
    client,database,_=chat;h=sign_in(client);base,url=begin(client,h)
    a,p,source=material(client,h,base,database);v=client.get(url).json();assert v['materials'][0]['reading_status']=='readable'
    row={'id':str(uuid4()),'category':'disbursement','content':'材料记载转账10万元，性质有争议','event_date':'2026-01-02','review_state':'pending','note':'与另一说法并列','source':source}
    manual=deepcopy(row);manual.update(id=str(uuid4()),category='other',content='人工核对记录，不证明其真实性')
    manual['source']={'kind':'manual_excerpt','attachment_id':a['id'],'original_sha256':a['sha256'],'location':'第1页人工摘录','quote':'甲向乙转账100000元'}
    body=request(v,[row,manual]);r=client.post(url+'/commits',headers=h,json=body);assert r.status_code==200,r.text
    # New material selection cannot rewrite the old quote or lose saved facts.
    scope=client.get(base).json()['scope'];assert client.put(base+'/materials',headers=h,json={'request_id':str(uuid4()),'expected_scope':scope,'items':[]}).status_code==200
    v=client.get(url).json();assert set(v['historical_ids'])=={row['id'],manual['id']}
    row['review_state']='disputed';r=client.post(url+'/commits',headers=h,json=request(v,[row,manual]));assert r.status_code==200,r.text
    # Old source may be retained for an existing row, not minted onto a new row.
    extra=deepcopy(row);extra['id']=str(uuid4());assert client.post(url+'/commits',headers=h,json=request(client.get(url).json(),[row,manual,extra])).status_code==409


@pytest.mark.parametrize('fault',['scope','profile','base','quote','cross','duplicate'])
def test_stale_or_forged_source_is_rejected_without_writing(chat,fault):
    client,database,_=chat;h=sign_in(client);base,url=begin(client,h)
    assert client.get(url).status_code==200
    row=statement(client,h,base,url);body=request(client.get(url).json(),[row])
    if fault=='scope':body['expected_scope']['input_revision']+=1
    if fault=='profile':body['expected_profile_revision']+=1
    if fault=='base':body['base_revision']+=1
    if fault=='quote':body['rows'][0]['source'].update(quote='伪造原话',end=4)
    if fault=='cross':body['expected_scope']['conversation_id']=str(uuid4())
    if fault=='duplicate':body['rows'].append(row)
    assert client.post(url+'/commits',headers=h,json=body).status_code in (409,422)
    with database.session() as db:assert db.scalar(text('SELECT count(*) FROM assistant_evidence_revisions'))==0


def test_document_tampering_and_cross_conversation(chat):
    client,database,_=chat;h=sign_in(client);base,url=begin(client,h)
    assert client.get(url).status_code==200
    a,p,source=material(client,h,base,database)
    row={'id':str(uuid4()),'category':'repayment','content':'还款待核','event_date':None,'review_state':'pending','note':'','source':source}
    for field,val in [('original_sha256','a'*64),('result_sha256','b'*64),('attachment_id',str(uuid4()))]:
        bad=deepcopy(row);bad['source'][field]=val
        assert client.post(url+'/commits',headers=h,json=request(client.get(url).json(),[bad])).status_code in (403,409,422)
    bad=deepcopy(row);bad['source']['anchor']['element_id']='伪造位置'
    assert client.post(url+'/commits',headers=h,json=request(client.get(url).json(),[bad])).status_code==409
    other,other_url=begin(client,h)
    assert client.post(other_url+'/commits',headers=h,json=request(client.get(other_url).json(),[row])).status_code in (403,409,422)


def test_owner_csrf_and_after_body_revocation(chat,monkeypatch):
    from cn_contract.web import evidence_worksheet
    client,database,_=chat;h=sign_in(client);base,url=begin(client,h)
    assert client.get(url).status_code==200
    row=statement(client,h,base,url);body=request(client.get(url).json(),[row]);original=evidence_worksheet.body_input
    async def reject(request,model):raise AssertionError('must authorize before reading body')
    monkeypatch.setattr(evidence_worksheet,'body_input',reject)
    assert client.post(url+'/commits',json=body).status_code==403
    sign_in(client,'另一合成律师');assert client.get(url).status_code==403
    h=sign_in(client)
    async def revoke(request,model):
        value=await original(request,model)
        with database.session() as db:db.execute(text('UPDATE sessions SET revoked=1'))
        return value
    monkeypatch.setattr(evidence_worksheet,'body_input',revoke)
    assert client.post(url+'/commits',headers=h,json=body).status_code in (401,403)


def test_actual_backup_restores_facts_receipts_and_original_readonly(native,tmp_path):
    from cn_contract.ops import workspace_backup
    from cn_contract.web.local_server import build_local_app
    from tests.helpers.legacy_assistant_backup import PASSPHRASE
    app,root=native
    with TestClient(app) as client:
        h=sign_in(client);base,url=begin(client,h);assert client.get(url).status_code==200
        a,p,source=material(client,h,base,app.state.database)
        row={'id':str(uuid4()),'category':'disbursement','content':'合成转账，性质仍待核','event_date':None,'review_state':'reviewed','note':'仅表示律师登记核对','source':source}
        body=request(client.get(url).json(),[row]);receipt=client.post(url+'/commits',headers=h,json=body).json()
    archive=tmp_path/'合成事实底稿.qxbak';workspace_backup.create_workspace_backup(root,archive,PASSPHRASE,data_parent=tmp_path)
    restored=workspace_backup.restore_workspace_backup(archive,tmp_path/'restored',PASSPHRASE);rr=Path(restored['restored_root']);app=build_local_app(rr,data_parent=rr.parent.parent)
    try:
        with TestClient(app) as client:
            h=sign_in(client);v=client.get(url).json();assert v['rows']==[row] and v['writable'] is False
            assert client.get(url+'/receipts/'+body['request_id']).json()==receipt
            assert client.get(url+'/revisions/1').json()['rows']==[row]
            assert client.get(base+'/attachments/'+a['id']+'/original').status_code==200
            assert client.post(url+'/commits',headers=h,json=body).status_code==409
            assert client.get(url+'/unsafe').status_code==409
    finally:app.state.database.engine.dispose()


def test_png_original_can_be_manually_excerpted_without_claiming_parsed_text(chat):
    import base64
    client,database,_=chat;h=sign_in(client);base,url=begin(client,h)
    raw=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jN1UAAAAASUVORK5CYII=')
    database._assistant_synthetic_hashes=frozenset({sha256(raw).hexdigest()})
    body={'request_id':str(uuid4()),'expected_scope':client.get(base).json()['scope'],
        'original_name':'合成空白凭证.png','mime':'image/png',
        'policy':{'policy_id':'assistant-local-files-v2','local_processing_accepted':True}}
    r=client.post(base+'/attachments',headers=h,files=[('file',(body['original_name'],raw,'image/png')),('metadata',(None,json.dumps(body)))])
    assert r.status_code==201,r.text
    a=r.json()['attachment'];v=client.get(url).json()
    assert v['materials'][0]['reading_status'] is None and v['materials'][0]['parse_revision'] is None
    # This deliberately fictional annotation proves manual provenance, not OCR accuracy.
    row={'id':str(uuid4()),'category':'other','content':'人工示例注记，空白像素不证明任何款项','event_date':None,
        'review_state':'pending','note':'仅测试原件关联，不声称图中有文字',
        'source':{'kind':'manual_excerpt','attachment_id':a['id'],'original_sha256':a['sha256'],'location':'1像素合成原件','quote':'人工示例注记'}}
    assert client.post(url+'/commits',headers=h,json=request(v,[row])).status_code==200
    assert client.get(url).json()['rows']==[row]
    assert client.get(base+'/attachments/'+a['id']+'/original').content==raw
    with database.session() as db:assert db.scalar(text('SELECT count(*) FROM assistant_attachment_parses'))==0
