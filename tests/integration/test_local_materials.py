"""Local-only workpapers use independent fiction outside every admitted registry."""
from io import BytesIO
import time
from uuid import UUID

from docx import Document
from fastapi.testclient import TestClient
import pytest

from cn_contract.documents.ingest import DOCX_MIME
from tests.integration.test_local_workbench import local, login, PASSWORD
from tests.integration.test_assisted_review import assisted_api


def new_local_case(client, headers):
    response = client.post('/api/matters', headers=headers, json={
        'title': '独立虚构本机材料验收', 'contract_type': 'sales',
        'review_mode': 'pre_sign', 'position': 'buyer',
        'objectives': ['核对本份虚构材料'], 'material_mode': 'local'})
    assert response.status_code == 201, response.text
    return response.json()['id']


def independent_docx():
    document = Document()
    document.add_paragraph('独立虚构材料，仅供工程测试。')
    document.add_paragraph('虚构甲方向虚构乙方采购测试设备，总价人民币12345元。')
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def upload_and_parse(client, headers, matter, raw):
    sent = client.post(f'/api/matters/{matter}/documents', headers={**headers, 'Idempotency-Key': 'local-upload'},
        files={'file': ('independent.docx', raw, DOCX_MIME)},
        data={'baseline_view': 'original', 'local_processing_confirmed': 'yes', 'local_policy_id': 'local-workpaper-v1'})
    assert sent.status_code == 202, sent.text
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        status = client.get('/api/jobs/' + sent.json()['job_id']).json()
        if status['state'] in ('completed', 'failed', 'blocked'):
            break
        time.sleep(.02)
    assert status['state'] == 'completed', status
    return sent.json()


def ready_local(client, headers, matter, sent):
    base = f'/api/matters/{matter}'
    parsed = client.get(base + f'/documents/{sent["version_id"]}/anchors').json()
    anchor = next(a for a in parsed['anchors'] if '12345' in a['quote'])
    facts = client.get(base + '/facts').json()
    response = client.post(base + '/facts/candidates', headers=headers, json={
        'document_version_id': sent['version_id'], 'document_sha256': sent['sha256'],
        'fact_version': facts['version'], 'candidates': [{'field': 'total_price',
            'typed_value': {'kind': 'decimal', 'value': '12345'}, 'unit': '元', 'anchor_ids': [anchor['id']]}]})
    assert response.status_code == 201, response.text
    candidate = response.json()['facts'][0]
    confirmed = client.post(base + '/facts/confirm', headers=headers, json={
        'fact_version': response.json()['version'], 'decisions': [{'fact_id': candidate['id'],
            'decision': 'confirmed', 'confirmed_value': {'kind': 'decimal', 'value': '12345'},
            'reason': '已核对独立虚构原文中的测试金额'}]})
    assert confirmed.status_code == 200, confirmed.text
    context = client.get(base + '/review-context')
    assert context.status_code == 200, context.text
    assert context.json()['candidate_source'] == 'manual_local'
    return anchor, confirmed.json(), context.json()


def test_local_manual_opinion_to_word_has_honest_source_and_old_draft_invalidates(local):
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client)
        matter = new_local_case(client, headers)
        base = f'/api/matters/{matter}'
        sent = upload_and_parse(client, headers, matter, independent_docx())
        anchor, facts, context = ready_local(client, headers, matter, sent)
        opinion = {'coverage_ids': ['C06'], 'issue_type': 'manual_price_scope', 'severity': 'low',
            'evidence_state': 'unconfirmed', 'anchor_ids': [anchor['id']],
            'analysis': '测试候选：价款范围需核对。', 'business_effect': '影响付款安排。',
            'recommendation': '明确价款包含项目。', 'basis_kind': 'material_observation',
            'authority_refs': [], 'unknowns': ['履行凭据待核。']}
        saved = client.post(base + '/review-proposals', headers={**headers, 'Idempotency-Key': 'manual-local'},
            json={'context_sha256': context['context_sha256'], 'fact_version': facts['version'],
                'document_hashes': {sent['version_id']: sent['sha256']},
                'examined_anchor_ids': [anchor['id']], 'opinions': [opinion]})
        assert saved.status_code == 201, saved.text
        run = client.get(base + '/reviews/' + saved.json()['id']).json()
        assert run['proposal']['execution_receipt']['source'] == 'manual_local'
        assert run['model_usage'] == []
        finding = run['findings'][0]
        decision = {'run_id': run['id'], 'run_sha256': run['run_sha256'], 'decision': 'accept', 'reason': '已核对本份测试意见'}
        response = client.post(base + f'/findings/{finding["id"]}/decision', headers=headers, json=decision)
        assert response.status_code == 201, response.text
        preview = client.post(base + '/exports/preview', headers=headers, json={
            'run_id': run['id'], 'run_sha256': run['run_sha256'], 'audience': 'client', 'patch_ids': []})
        assert preview.status_code == 201, preview.text
        artifact = preview.json()['artifacts'][0]
        url = base + f'/exports/{artifact["id"]}/download'
        response = client.get(url)
        assert response.status_code == 200, response.text
        document = Document(BytesIO(response.content))
        assert document.core_properties.keywords == 'DRAFT MANUAL-LOCAL'
        assert '虚构' not in document.core_properties.title + document.core_properties.subject
        assert '本机人工' in document.sections[0].header.paragraphs[0].text
        text = '\n'.join([p.text for p in document.paragraphs] +
            [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        assert opinion['analysis'] in text and opinion['recommendation'] in text
        assert '离线人工/合成' not in text
        assert '模拟' not in text
        for audience in ('internal', 'counterparty'):
            exported = client.post(base + '/exports/preview', headers=headers, json={
                'run_id': run['id'], 'run_sha256': run['run_sha256'], 'audience': audience, 'patch_ids': []})
            assert exported.status_code == 201, exported.text
            payload = client.get(base + '/exports/' + exported.json()['artifacts'][0]['id'] + '/download')
            assert payload.status_code == 200
            report = Document(BytesIO(payload.content))
            assert all('模拟' not in p.text and '插件候选' not in p.text for p in report.paragraphs)
        approved = client.post(base + '/patches/approve', headers=headers, json={
            'run_id': run['id'], 'run_sha256': run['run_sha256'], 'finding_id': finding['id'],
            'anchor_id': anchor['id'], 'new_text': '虚构甲方向虚构乙方采购测试设备，总价人民币12345元，价款包含安装。',
            'reason': '测试律师逐段批准', 'audiences': ['internal']})
        assert approved.status_code == 201, approved.text
        internal = client.post(base + '/exports/preview', headers=headers, json={
            'run_id': run['id'], 'run_sha256': run['run_sha256'], 'audience': 'internal', 'patch_ids': [approved.json()['id']]})
        assert internal.status_code == 201, internal.text
        revised = next(a for a in internal.json()['artifacts'] if a['type'] == 'redline_contract_docx')
        redline = client.get(base + f'/exports/{revised["id"]}/download')
        assert redline.status_code == 200
        from cn_contract.exports.redline_adapter import baseline_texts
        assert baseline_texts(redline.content, accept=False) == [p.text for p in Document(BytesIO(independent_docx())).paragraphs]
        assert '价款包含安装' in baseline_texts(redline.content, accept=True)[1]
        assert '模拟' not in Document(BytesIO(redline.content)).sections[0].header.paragraphs[0].text
        decision['decision'] = 'reject'
        assert client.post(base + f'/findings/{finding["id"]}/decision', headers=headers, json=decision).status_code == 201
        assert client.get(url).status_code == 409


@pytest.mark.parametrize('form,mime', [
    ({}, DOCX_MIME),
    ({'local_processing_confirmed': 'no', 'local_policy_id': 'local-workpaper-v1'}, DOCX_MIME),
    ({'local_processing_confirmed': 'yes', 'local_policy_id': 'wrong-policy'}, DOCX_MIME),
    ({'local_processing_confirmed': 'yes', 'local_policy_id': 'local-workpaper-v1'}, 'application/pdf')])
def test_missing_or_wrong_local_consent_and_pdf_leave_no_original_or_document(local, form, mime):
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client); matter = new_local_case(client, headers)
        response = client.post(f'/api/matters/{matter}/documents', headers={**headers, 'Idempotency-Key': 'denied'},
            files={'file': ('independent.docx', independent_docx(), mime)}, data={'baseline_view': 'original', **form})
        assert response.status_code == 409, response.text
        assert client.get(f'/api/matters/{matter}').json()['documents'] == []
        assert not list((root / 'protected').rglob('original.bin'))


def test_local_case_cannot_share_with_plugin_or_use_simulated_queries(local):
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client); matter = new_local_case(client, headers)
        sent = upload_and_parse(client, headers, matter, independent_docx())
        _, _, context = ready_local(client, headers, matter, sent)
        response = client.post(f'/api/matters/{matter}/case-bridge', headers=headers,
            json={'context_sha256': context['context_sha256'], 'share_with_chat': True})
        assert response.status_code == 409 and response.json()['code'] == 'DATA_POLICY_UNAPPROVED'
        assert not (root / 'selected-case.json').exists()
        response = client.get(f'/api/matters/{matter}/query-status')
        assert response.status_code == 409 and response.json()['code'] == 'DATA_POLICY_UNAPPROVED'


def test_local_case_with_no_original_is_excluded_from_synthetic_backup(local, tmp_path):
    from cn_contract.domain import BackupConfig, Role
    from cn_contract.security import authenticate
    from cn_contract.ops.backup import backup_and_verify
    from cn_contract.security import ServiceError
    from tests.helpers.synthetic_operations import keys
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    module.add_local_user(root, 'local-admin', PASSWORD, Role.ADMIN)
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client); new_local_case(client, headers)
        login(client, 'local-admin')
        database = client.app.state.database
        with database.session() as db:
            actor = authenticate(client.cookies.get('qx_session'), db)
        with pytest.raises(ServiceError) as denied:
            backup_and_verify(BackupConfig(storage_ref=str(tmp_path/'archives'), key_ref='synthetic-key-1',retention_days=30),
                database=database, actor=actor, key_resolver=keys(), restore_parent=tmp_path/'restores')
        assert denied.value.code == 'DATA_POLICY_UNAPPROVED'
        assert not (tmp_path/'archives').exists()


def test_local_forms_require_file_consent_and_hide_plugin_action(local):
    from lxml import html
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client)
        cases = html.fromstring(client.get('/').text)
        assert cases.xpath('//select[@name="material_mode"]/option[@value="local"]')
        matter = new_local_case(client, headers)
        intake = html.fromstring(client.get('/matters/' + matter).text)
        assert intake.xpath('//input[@name="local_processing_confirmed"][@required][@value="yes"]')
        assert intake.xpath('//input[@name="local_policy_id"][@value="local-workpaper-v1"]')
        review = html.fromstring(client.get('/matters/' + matter + '/review').text)
        assert not review.xpath('//form[@id="case-bridge-form"]')
        exports = html.fromstring(client.get('/matters/' + matter + '/exports').text)
        assert exports.xpath('//form[@id="patch-form"]//input[@name="reason"]/@value') == ['']


def test_demo_server_cannot_enable_local_case_from_http(assisted_api):
    from tests.integration.test_api import sign_in
    client, _, _, _ = assisted_api
    headers = sign_in(client)
    response = client.post('/api/matters', headers=headers, json={
        'title': '虚构阻断测试', 'contract_type': 'sales', 'review_mode': 'pre_sign',
        'position': 'buyer', 'objectives': ['虚构验证'], 'material_mode': 'local'})
    assert response.status_code == 409 and response.json()['code'] == 'DATA_POLICY_UNAPPROVED'
    assert client.get('/api/matters').json() == []


def test_local_file_confirmation_cannot_transfer_to_another_case(local):
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client)
        first = new_local_case(client, headers); raw = independent_docx()
        sent = upload_and_parse(client, headers, first, raw)
        second = new_local_case(client, headers)
        response = client.post(f'/api/matters/{second}/documents', headers={**headers, 'Idempotency-Key': 'second'},
            files={'file': ('same.docx', raw, DOCX_MIME)}, data={'baseline_view': 'original'})
        assert response.status_code == 409
        assert client.get(f'/api/matters/{second}').json()['documents'] == []
        assert client.get(f'/api/matters/{second}/documents/{sent["version_id"]}/anchors').status_code == 403


def test_local_file_replay_new_version_and_restart_keep_independent_confirmation(local):
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    raw = independent_docx()
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client); matter = new_local_case(client, headers)
        sent = upload_and_parse(client, headers, matter, raw)
        url = f'/api/matters/{matter}/documents'
        files = {'file': ('independent.docx', raw, DOCX_MIME)}
        replay = client.post(url, headers={**headers, 'Idempotency-Key': 'local-upload'}, files=files,
            data={'baseline_view': 'original', 'local_processing_confirmed': 'yes', 'local_policy_id': 'local-workpaper-v1'})
        assert replay.status_code == 202 and replay.json() == sent
        denied = client.post(url, headers={**headers, 'Idempotency-Key': 'v2'}, files=files,
            data={'baseline_view': 'original', 'supersedes_id': sent['version_id']})
        assert denied.status_code == 409
        newer = client.post(url, headers={**headers, 'Idempotency-Key': 'v2'}, files=files,
            data={'baseline_view': 'original', 'supersedes_id': sent['version_id'],
                'local_processing_confirmed': 'yes', 'local_policy_id': 'local-workpaper-v1'})
        assert newer.status_code == 202, newer.text
        new_id = newer.json()['version_id']
        assert new_id != sent['version_id']
        assert client.get(url + '/' + sent['version_id'] + '/anchors').status_code == 409
    with TestClient(module.build_local_app(root)) as resumed:
        login(resumed)
        state = resumed.get(f'/api/matters/{matter}').json()
        assert state['synthetic'] is False and [v['id'] for v in state['documents']] == [new_id]
        from cn_contract.persistence import read_original
        from cn_contract.security import authenticate
        with resumed.app.state.database.session() as db:
            actor = authenticate(resumed.cookies.get('qx_session'), db)
            assert read_original(UUID(matter), UUID(sent['version_id']), actor, db) == raw
            assert read_original(UUID(matter), UUID(new_id), actor, db) == raw


def test_failed_local_upload_transaction_preserves_no_original_or_admission(local, monkeypatch):
    from cn_contract.web import storage
    from cn_contract.persistence import AuditRow
    from sqlalchemy import select
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client); matter = new_local_case(client, headers)
        def fail(*args, **kwargs):
            raise RuntimeError('independent simulated queue persistence failure')
        monkeypatch.setattr(storage, 'enqueue', fail)
        with pytest.raises(RuntimeError, match='independent simulated queue'):
            client.post(f'/api/matters/{matter}/documents', headers={**headers, 'Idempotency-Key': 'fail'},
                files={'file': ('independent.docx', independent_docx(), DOCX_MIME)},
                data={'baseline_view': 'original', 'local_processing_confirmed': 'yes', 'local_policy_id': 'local-workpaper-v1'})
        assert client.get(f'/api/matters/{matter}').json()['documents'] == []
        assert not list((root/'protected').rglob('original.bin'))
        with client.app.state.database.session() as db:
            assert not list(db.scalars(select(AuditRow).where(AuditRow.action == 'local_material_admitted')))


def test_explicit_local_case_import_preserves_original_and_parses_without_mock(local):
    module, root = local
    module.initialize_workspace(root, 'local-lawyer', PASSWORD)
    raw = independent_docx()
    with TestClient(module.build_local_app(root)) as client:
        headers = login(client)
        matter = new_local_case(client, headers)
        sent = client.post(f'/api/matters/{matter}/documents', headers={**headers, 'Idempotency-Key': 'local-upload'},
            files={'file': ('independent.docx', raw, DOCX_MIME)},
            data={'baseline_view': 'original', 'local_processing_confirmed': 'yes',
                  'local_policy_id': 'local-workpaper-v1'})
        assert sent.status_code == 202, sent.text
        assert client.get(f'/api/matters/{matter}').json()['synthetic'] is False
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            status = client.get('/api/jobs/' + sent.json()['job_id']).json()
            if status['state'] in ('completed', 'failed'):
                break
            time.sleep(.02)
        assert status['state'] == 'completed', status
        facts = client.get(f'/api/matters/{matter}/facts').json()
        assert facts['facts'] == []
        from cn_contract.persistence import read_original
        from cn_contract.security import authenticate
        with client.app.state.database.session() as db:
            actor = authenticate(client.cookies.get('qx_session'), db)
            assert read_original(UUID(matter), UUID(sent.json()['version_id']), actor, db) == raw
