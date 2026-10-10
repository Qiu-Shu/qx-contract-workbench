"""Fictional OOXML visibility and independent lawyer-source boundaries."""
import hashlib
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.shared import RGBColor
from cn_contract.facts.storage import source_anchor_is_eligible
from cn_contract.documents.ingest import DOCX_MIME
from tests.integration.test_docx_office_parts import saved, parse
from tests.integration.test_assisted_review import assisted_api
from tests.integration.test_api import sign_in, create_case
import pytest


def colored_document(kind='direct', foreground='FFFFFF', background=None):
    document = Document()
    p = document.add_paragraph('SYNTHETIC Price: ')
    r = p.add_run('1')
    if kind in ('paragraph', 'character'):
        style = document.styles.add_style('SyntheticColor', WD_STYLE_TYPE.PARAGRAPH if kind == 'paragraph' else WD_STYLE_TYPE.CHARACTER)
        style.font.color.rgb = RGBColor.from_string(foreground)
        child = document.styles.add_style('SyntheticDerived', style.type)
        child.base_style = style
        if kind == 'paragraph': p.style = child
        else: r.style = child
    else:
        r.font.color.rgb = RGBColor.from_string(foreground)
        if kind == 'theme':
            r._r.rPr.color.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}themeColor', 'accent1')
    if background:
        shade = OxmlElement('w:shd')
        shade.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill', background)
        r._r.get_or_add_rPr().append(shade)
    p.add_run('28473 CNY')
    document.add_paragraph('SYNTHETIC ordinary visible paragraph')
    return saved(document)


@pytest.mark.parametrize('kind,foreground,background,risky', [
    ('direct','FFFFFF',None,True), ('paragraph','FFFFFF',None,True),
    ('character','FFFFFF',None,True), ('theme','000000',None,True),
    ('direct','EEEEEE',None,True), ('direct','FFFFFF','000000',False),
    ('direct','000000',None,False)])
def test_color_risk_preserves_text_but_controls_eligibility(kind, foreground, background, risky):
    result = parse(colored_document(kind,foreground,background))
    amount = next(a for a in result.anchors if '128473' in a.quote)
    assert source_anchor_is_eligible(amount,result) is not risky
    assert bool(result.quality_issues) is risky
    assert source_anchor_is_eligible(result.anchors[-1],result)


def setup_colored(assisted_api):
    raw=colored_document(); digest=hashlib.sha256(raw).hexdigest()
    client,database,ids,module=assisted_api.setup(extra_hashes=(digest,))
    headers=sign_in(client); matter=create_case(client,headers,title='独立虚构DOCX来源核对')
    sent=client.post(f'/api/matters/{matter}/documents',headers={**headers,'Idempotency-Key':'color'},
        files={'file':('synthetic.docx',raw,DOCX_MIME)},data={'baseline_view':'original'}).json()
    assert module.run_next_job(database).state=='completed'
    parsed=client.get(f'/api/matters/{matter}/documents/{sent["version_id"]}/anchors').json()
    anchor=next(a for a in parsed['anchors'] if '128473' in a['quote'])
    facts=client.get(f'/api/matters/{matter}/facts').json()
    return client,headers,matter,sent,digest,anchor,facts,raw


def test_default_table_style_color_is_not_assumed_visible():
    document=Document()
    style=document.styles['Normal Table']
    style.font.color.rgb=RGBColor(255,255,255)
    document.add_table(rows=1,cols=1).cell(0,0).text='SYNTHETIC 128473 CNY'
    result=parse(saved(document))
    assert result.anchors and not source_anchor_is_eligible(result.anchors[0],result)


def test_direct_paragraph_background_overrides_inherited_dark_background():
    from docx.oxml.ns import qn
    document=Document()
    style=document.styles.add_style('SyntheticDark',WD_STYLE_TYPE.PARAGRAPH)
    def shade(properties, fill):
        node=OxmlElement('w:shd');node.set(qn('w:fill'),fill);properties.append(node)
    shade(style.element.get_or_add_pPr(),'000000')
    p=document.add_paragraph(style=style)
    shade(p._p.get_or_add_pPr(),'FFFFFF')
    p.add_run('1').font.color.rgb=RGBColor(255,255,255)
    r=p.add_run('28473 CNY');r.font.color.rgb=RGBColor(0,0,0)
    shade(r._r.get_or_add_rPr(),'FFFFFF')
    result=parse(saved(document))
    assert result.anchors[0].quote=='128473 CNY'
    assert not source_anchor_is_eligible(result.anchors[0],result)


def test_color_candidate_is_blocked_and_independent_source_can_continue(assisted_api):
    client,headers,matter,sent,digest,anchor,facts,raw=setup_colored(assisted_api)
    candidate=client.post(f'/api/matters/{matter}/facts/candidates',headers=headers,json={
        'document_version_id':sent['version_id'],'document_sha256':digest,'fact_version':facts['version'],
        'candidates':[{'field':'total_price','typed_value':{'kind':'decimal','value':'128473'},
                       'unit':'CNY','anchor_ids':[anchor['id']]}]})
    assert candidate.status_code==409 and candidate.json()['code']=='SOURCE_VISUAL_CONFIRMATION_REQUIRED'
    original=client.get(f'/api/matters/{matter}/documents/{sent["version_id"]}/original?sha256={digest}')
    assert original.status_code==200 and original.content==raw
    value={'document_version_id':sent['version_id'],'document_sha256':digest,'fact_version':facts['version'],
        'page':None,'source_anchor_id':anchor['id'],'source_kind':'lawyer_transcription',
        'observed_text':'SYNTHETIC Price: 28473 CNY','field':'total_price',
        'typed_value':{'kind':'decimal','value':'28473'},'unit':'CNY','reason':'独立虚构原件可见金额核对',
        'original_page_checked':True,'previous_id':None,'candidate_id':None,'candidate_sha256':None}
    response=client.post(f'/api/matters/{matter}/page-verifications',headers={**headers,'Idempotency-Key':'source'},json=value)
    assert response.status_code==200,response.text
    snapshot=client.get(f'/api/matters/{matter}/facts').json()
    fact=snapshot['facts'][-1]
    assert fact['anchor_ids']==[] and fact['confirmation_status']=='unconfirmed'
    confirmed=client.post(f'/api/matters/{matter}/facts/confirm',headers=headers,json={
        'fact_version':snapshot['version'],'decisions':[{'fact_id':fact['id'],'decision':'confirmed',
        'confirmed_value':fact['typed_value'],'reason':'独立核对后确认'}]})
    assert confirmed.status_code==200,confirmed.text
    context=client.get(f'/api/matters/{matter}/review-context')
    assert context.status_code==200,context.text
    assert anchor['id'] not in context.json()['eligible_anchor_ids']
    assert response.json()['id'] in context.json()['eligible_verification_ids']
    ctx=context.json(); record_id=response.json()['id']
    proposal={'context_sha256':ctx['context_sha256'],'fact_version':confirmed.json()['version'],
        'document_hashes':{sent['version_id']:digest},'examined_anchor_ids':[],
        'examined_verification_ids':[record_id],'opinions':[{
        'coverage_ids':['C06'],'issue_type':'synthetic_source','severity':'low','evidence_state':'unconfirmed',
        'anchor_ids':[],'verification_ids':[record_id],'analysis':'虚构核对金额28473 CNY',
        'business_effect':'来源须保留','recommendation':'以人工核对记录为准',
        'basis_kind':'material_observation','authority_refs':[],'unknowns':[]}]}
    review=client.post(f'/api/matters/{matter}/review-proposals',headers={**headers,'Idempotency-Key':'review'},json=proposal)
    assert review.status_code==201,review.text
    run=client.get(f'/api/matters/{matter}/reviews/{review.json()["id"]}').json()
    decision=client.post(f'/api/matters/{matter}/findings/{run["findings"][0]["id"]}/decision',headers=headers,
        json={'run_id':run['id'],'run_sha256':run['run_sha256'],'decision':'accept','reason':'独立虚构核对'})
    assert decision.status_code==201,decision.text
    bundle=client.post(f'/api/matters/{matter}/exports/preview',headers=headers,
        json={'run_id':run['id'],'run_sha256':run['run_sha256'],'audience':'client','patch_ids':[]})
    assert bundle.status_code==201,bundle.text
    download=client.get(f'/api/matters/{matter}/exports/{bundle.json()["artifacts"][0]["id"]}/download')
    from io import BytesIO
    text='\n'.join(p.text for p in Document(BytesIO(download.content)).paragraphs)
    assert '28473' in text and 'DOCX段落' in text and '第 None 页' not in text


@pytest.mark.parametrize('change', ['anchor','hash','page','unchecked','candidate'])
def test_docx_source_record_rejects_unbound_or_unchecked_source(assisted_api,change):
    from uuid import uuid4
    client,headers,matter,sent,digest,anchor,facts,_=setup_colored(assisted_api)
    value={'document_version_id':sent['version_id'],'document_sha256':digest,'fact_version':facts['version'],
        'page':None,'source_anchor_id':anchor['id'],'source_kind':'lawyer_transcription',
        'observed_text':'SYNTHETIC Price: 28473 CNY','field':'total_price',
        'typed_value':{'kind':'decimal','value':'28473'},'unit':'CNY','reason':'独立虚构核对',
        'original_page_checked':True,'previous_id':None,'candidate_id':None,'candidate_sha256':None}
    if change=='anchor': value['source_anchor_id']=str(uuid4())
    elif change=='hash': value['document_sha256']='0'*64
    elif change=='page': value['page']=1
    elif change=='unchecked': value['original_page_checked']=False
    else: value.update(candidate_id=str(uuid4()),candidate_sha256='0'*64)
    response=client.post(f'/api/matters/{matter}/page-verifications',headers={**headers,'Idempotency-Key':'bad'},json=value)
    assert response.status_code in (409,422),response.text
    assert client.get(f'/api/matters/{matter}/facts').json()==facts
