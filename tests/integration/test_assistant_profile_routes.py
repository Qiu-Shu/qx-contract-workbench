"""Authenticated manual profile transport over independent synthetic local workspaces."""
from hashlib import sha256
from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
import pytest
from sqlalchemy import text

from test_assistant_routes import native, chat, create, sign_in, sent


def command(client, headers, *, project_intake=True):
    first = create(client, headers)
    sent(client, headers, first)
    url = '/api/assistant/conversations/' + first['scope']['conversation_id']
    current = client.get(url).json()
    message = current['messages'][0]
    ref = {'kind':'message', 'conversation_id':current['scope']['conversation_id'],
           'message_id':message['id'], 'source_revision':message['source_revision'],
           'start':0, 'end':len(message['content']), 'quote':message['content'],
           'quote_sha256':sha256(message['content'].encode()).hexdigest()}
    return url + '/profile', {'request_id':str(uuid4()), 'patch':{'scope':current['scope'],
        'base_profile_revision':0, 'changes':[{'field':'consultation', 'origin':'model_inference',
        'value':{'kind':'text','value':'合成接待摘要，客户称尚未收款，待核验。'}, 'source_refs':[ref]}]},
        'project_intake':project_intake}


def commit(client, headers, url, value):
    result = client.post(url+'/commits', headers=headers, json=value)
    assert result.status_code == 200, result.text
    return result.json()


def test_empty_profile_route_exists(chat):
    client, _, _ = chat
    headers = sign_in(client)
    first = create(client, headers)
    url = '/api/assistant/conversations/'+first['scope']['conversation_id']+'/profile'
    response = client.get(url)
    assert response.status_code == 200, response.text
    assert response.json() is None
    assert client.get(url+'/receipts/'+str(uuid4())).json() is None


def test_commit_history_receipt_and_intake_source_label(chat):
    client, database, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    receipt = commit(client, headers, url, value)
    assert receipt['audience'] == 'internal'
    assert receipt['event_type'] == 'profile.saved'
    assert receipt['consultation_origin'] == 'profile_summary'
    assert commit(client, headers, url, value) == receipt
    assert client.get(url+'/receipts/'+value['request_id']).json() == receipt
    current = client.get(url).json()
    assert current['revision'] == 1
    assert client.get(url+'/revisions/1').json() == current
    assert current['fields'][0]['verification'] == 'unverified'
    intake_url = '/api/intakes/'+receipt['intake_id']
    intake = client.get(intake_url).json()
    assert intake['consultation_origin'] == 'profile_summary'
    assert intake['source_conversation_id'] == receipt['conversation_id']
    assert intake['source_profile_revision'] == 1
    page = client.get('/intakes/'+receipt['intake_id'])
    assert page.status_code == 200 and '接待摘要' in page.text
    assert '/assistant/'+receipt['conversation_id'] in page.text
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_profile_receipts')) == 1
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_model_attempts')) == 0
        assert db.scalar(text('SELECT COUNT(*) FROM matters')) == 0


def test_explicit_manual_edits_are_protected_and_never_confirmed(chat):
    client, _, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers, project_intake=False)
    value['manual_edits'] = [{'field':'consultation','value':{'kind':'text','value':'律师更正的合成内容'}}]
    receipt = commit(client, headers, url, value)
    assert receipt['intake_id'] is None
    field = client.get(url).json()['fields'][0]
    assert field['protected_value']['value'] == '律师更正的合成内容'
    assert field['verification'] != 'confirmed'


@pytest.mark.parametrize('payload_change', [
    {'owner_id':str(uuid4())}, {'audience':'client'}, {'project_intake':'true'},
    {'expected_intake_id':str(uuid4())}, {'confirmed':True},
])
def test_unknown_or_inconsistent_commands_fail_closed(chat, payload_change):
    client, database, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    result = client.post(url+'/commits', headers=headers, json={**value,**payload_change})
    assert result.status_code == 422
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_profile_receipts')) == 0


def test_path_scope_mismatch_is_rejected_before_store(chat, monkeypatch):
    from cn_contract.assistant import profile_store
    client, _, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    other = create(client, headers)
    value['patch']['scope'] = other['scope']
    monkeypatch.setattr(profile_store, 'save_profile', lambda *args, **kwargs: pytest.fail('mismatched route reached store'))
    response = client.post(url+'/commits', headers=headers, json=value)
    assert response.status_code == 409


def test_ownership_and_authentication_precede_body_read(chat, monkeypatch):
    from cn_contract.web import assistant_profiles
    client, _, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    commit(client, headers, url, value)
    other = sign_in(client, '另一合成律师')
    async def unread(*args):
        pytest.fail('unauthorized body was read')
    monkeypatch.setattr(assistant_profiles, 'body_input', unread)
    assert client.post(url+'/commits', headers=other, content='invalid-json').status_code == 403
    for suffix in ('','/revisions/1','/receipts/'+value['request_id']):
        assert client.get(url+suffix).status_code == 403
    client.cookies.clear()
    assert client.post(url+'/commits', content='invalid-json').status_code == 401
    assert client.get(url).status_code == 401


def test_csrf_origin_body_limit_and_revision_boundaries(chat):
    client, _, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    assert client.post(url+'/commits', json=value).status_code == 403
    assert client.post(url+'/commits', headers={**headers,'Origin':'https://evil.invalid'},json=value).status_code == 403
    assert client.post(url+'/commits', headers=headers, content='{"request_id":"a","request_id":"b"}').status_code == 422
    assert client.post(url+'/commits', headers=headers, content='x'*(256*1024+1)).status_code == 413
    assert client.get(url+'/revisions/0').status_code == 422
    assert client.get(url+'/revisions/-1').status_code == 422


def test_revoked_session_after_body_parsing_is_rechecked(chat, monkeypatch):
    from cn_contract.web import assistant_profiles
    client, database, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    original = assistant_profiles.body_input
    async def revoke(request, model):
        value = await original(request, model)
        with database.session() as db:
            db.execute(text('UPDATE sessions SET revoked=1'))
        return value
    monkeypatch.setattr(assistant_profiles, 'body_input', revoke)
    assert client.post(url+'/commits', headers=headers, json=value).status_code == 401
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_profile_receipts')) == 0


def test_intake_export_and_manual_edit_keep_historical_origin(chat):
    from cn_contract import domain as d
    client, _, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    receipt = commit(client, headers, url, value)
    intake_url = '/api/intakes/'+receipt['intake_id']
    current = client.get(intake_url).json()
    exported = client.post(intake_url+'/export', headers=headers, json={'revision':1,'audience':'internal','text_checked':True})
    assert exported.status_code == 200, exported.text
    paragraphs = [p.text for p in Document(BytesIO(exported.content)).paragraphs]
    assert '接待摘要' in paragraphs and '咨询原话' not in paragraphs
    assert any(receipt['conversation_id'] in p for p in paragraphs)
    edited = {key:current[key] for key in d.IntakeDraft.model_fields}
    edited['consultation'] = '手工登记的新版本'
    result = client.put(intake_url, headers=headers, json={'expected_revision':1,'draft':edited})
    assert result.status_code == 200, result.text
    assert result.json()['consultation_origin'] == 'user'
    assert result.json()['source_conversation_id'] is None
    assert client.get(intake_url+'/revisions/1').json()['consultation_origin'] == 'profile_summary'
    exported = client.post(intake_url+'/export', headers=headers, json={'revision':2,'audience':'internal','text_checked':True})
    paragraphs = [p.text for p in Document(BytesIO(exported.content)).paragraphs]
    assert '咨询原话' in paragraphs and '接待摘要' not in paragraphs


def test_summary_source_survives_other_edits_but_not_consultation_replacement(chat):
    from cn_contract import domain as d
    client, _, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers)
    receipt = commit(client, headers, url, value)
    intake_url = '/api/intakes/'+receipt['intake_id']
    original = client.get(intake_url).json()
    edited = {key:original[key] for key in d.IntakeDraft.model_fields}
    edited['title'] = '只更改合成咨询标题'
    second = client.put(intake_url, headers=headers, json={'expected_revision':1,'draft':edited})
    assert second.status_code == 200
    assert second.json()['consultation_origin'] == 'profile_summary'
    assert second.json()['source_conversation_id'] == receipt['conversation_id']
    edited['consultation'] = '人工改写正文'
    third = client.put(intake_url, headers=headers, json={'expected_revision':2,'draft':edited})
    assert third.status_code == 200 and third.json()['consultation_origin'] == 'user'
    edited['consultation'] = original['consultation']
    fourth = client.put(intake_url, headers=headers, json={'expected_revision':3,'draft':edited})
    assert fourth.status_code == 200 and fourth.json()['consultation_origin'] == 'user'
    assert fourth.json()['source_conversation_id'] is None
    assert client.get(intake_url+'/revisions/2').json()['consultation_origin'] == 'profile_summary'


def test_recovery_reads_profile_history_and_summary_but_fences_writes(native):
    from cn_contract.ops import storage
    from cn_contract.web import local_server
    original, root = native
    database = original.state.database
    with TestClient(original) as client:
        headers = sign_in(client)
        url, value = command(client, headers)
        receipt = commit(client, headers, url, value)
    with database.session() as db:
        db.add(storage.RestoreFenceRow(ref=str(uuid4()),archive_sha256='a'*64,manifest_sha256='b'*64,payload='{}'))
    app = local_server.build_local_app(root, data_parent=root.parent.parent)
    with TestClient(app) as client:
        headers = sign_in(client)
        for suffix in ('','/revisions/1','/receipts/'+value['request_id']):
            assert client.get(url+suffix).status_code == 200
        assert client.post(url+'/commits', headers=headers, json=value).status_code == 409
        for suffix in ('/commits','/revisions/0','/revisions/1/unsafe','/receipts/'+value['request_id']+'/unsafe'):
            assert client.get(url+suffix).status_code == 409
        page = client.get('/recovery/intakes/'+receipt['intake_id'])
        assert page.status_code == 200 and '接待摘要' in page.text
        assert '/assistant/'+receipt['conversation_id'] in page.text


def test_manual_field_types_null_and_unchanged_receipts_survive_later_candidates(chat):
    client, database, _ = chat
    headers = sign_in(client)
    url, value = command(client, headers, project_intake=False)
    later_proposal = value['patch']['changes'][0]
    edits = [
        {'field':'title','value':{'kind':'null','value':None}},
        {'field':'contact','value':{'kind':'text','value':'合成律师维护的联系人'}},
        {'field':'consultation','value':{'kind':'text','value':'合成律师更正的咨询摘要'}},
        {'field':'questions','value':{'kind':'list','value':['合成追问甲','合成追问乙']}},
        {'field':'materials','value':{'kind':'materials','value':[{'name':'合成凭证','state':'checked'}]}},
        {'field':'follow_up_at','value':{'kind':'datetime','value':'2026-10-24T08:30:00+08:00'}},
        {'field':'follow_up_note','value':{'kind':'text','value':'合成跟进备注'}},
        {'field':'client_reply','value':{'kind':'text','value':'合成客户回复'}},
        {'field':'internal_notes','value':{'kind':'text','value':'合成内部意见'}},
        {'field':'principal_amount','value':{'kind':'money','value':'9999999999999999.12345678','unit':'CNY'}},
        {'field':'interest_amount','value':{'kind':'money','value':'0.00000001','unit':'CNY'}},
        {'field':'payment_amount','value':{'kind':'money','value':'0','unit':'USD'}},
    ]
    value['patch']['changes'] = []
    value['manual_edits'] = edits
    first = commit(client, headers, url, value)
    assert first['profile_revision'] == 1 and first['profile_changed']
    assert first['intake_id'] is None and not first['projection_changed']
    before = client.get(url).json()
    fields = {item['field']:item for item in before['fields']}
    assert len(fields) == 12
    assert fields['title']['protected_value'] == {'kind':'null','value':None}
    assert fields['contact']['protected_value']['value'] == '合成律师维护的联系人'
    assert fields['follow_up_at']['protected_value']['value'] == '2026-10-24T00:30:00Z'
    assert fields['principal_amount']['protected_value']['value'] == '9999999999999999.12345678'
    assert fields['materials']['protected_value']['value'][0]['state'] == 'checked'
    assert all(field['verification'] != 'confirmed' for field in fields.values())
    value['request_id'] = str(uuid4())
    value['patch']['base_profile_revision'] = 1
    unchanged = commit(client, headers, url, value)
    assert unchanged['event_type'] == 'profile.unchanged'
    assert not unchanged['profile_changed'] and unchanged['profile_revision'] == 1
    assert client.get(url+'/receipts/'+value['request_id']).json() == unchanged
    value['request_id'] = str(uuid4())
    value['manual_edits'] = []
    value['patch']['changes'] = [later_proposal]
    later = commit(client, headers, url, value)
    assert later['profile_revision'] == 2
    after = {item['field']:item for item in client.get(url).json()['fields']}
    assert {key:item['protected_value'] for key,item in after.items()} == {
        key:item['protected_value'] for key,item in fields.items()}
    assert after['consultation']['verification'] == 'conflicted'
    assert client.get(url+'/revisions/1').json() == before
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_model_attempts')) == 0
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_profile_receipts')) == 3


def test_manual_profile_edit_captures_intake_baseline_without_changing_intake(chat):
    client, database, _ = chat
    headers = sign_in(client)
    url, initial = command(client, headers)
    first = commit(client, headers, url, initial)
    intake_url = '/api/intakes/'+first['intake_id']
    before = client.get(intake_url).json()
    edit = {'request_id':str(uuid4()),'patch':{'scope':initial['patch']['scope'],
        'base_profile_revision':1,'changes':[]},'expected_intake_id':first['intake_id'],
        'expected_intake_revision':first['intake_revision'],
        'manual_edits':[{'field':'consultation','value':{'kind':'null','value':None}}],'project_intake':False}
    saved = commit(client, headers, url, edit)
    assert saved['profile_revision'] == 2 and saved['profile_changed']
    assert saved['intake_revision'] == before['revision'] and not saved['projection_changed']
    assert client.get(intake_url).json() == before
    later = {**edit,'request_id':str(uuid4()),'manual_edits':[],'project_intake':True,
        'patch':{**edit['patch'],'base_profile_revision':2,'changes':initial['patch']['changes']}}
    projected = commit(client, headers, url, later)
    assert client.get(intake_url).json()['consultation'] == ''
    fields = {item['field']:item for item in client.get(url).json()['fields']}
    assert fields['consultation']['protected_value'] == {'kind':'null','value':None}
    assert projected['consultation_origin'] == 'user'
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_model_attempts')) == 0


@pytest.mark.parametrize('stale', ['profile','intake','scope'])
def test_stale_manual_edit_baselines_do_not_write_or_create_receipts(chat, stale):
    client, database, _ = chat
    headers = sign_in(client)
    url, initial = command(client, headers)
    receipt = commit(client, headers, url, initial)
    before = client.get(url).json()
    value = {'request_id':str(uuid4()),'patch':{'scope':dict(initial['patch']['scope']),
        'base_profile_revision':receipt['profile_revision'],'changes':[]},
        'expected_intake_id':receipt['intake_id'],'expected_intake_revision':receipt['intake_revision'],
        'manual_edits':[{'field':'contact','value':{'kind':'text','value':'过期合成更正'}}],'project_intake':False}
    if stale == 'profile':
        value['patch']['base_profile_revision'] = 0
    elif stale == 'intake':
        value['expected_intake_revision'] += 1
    else:
        value['patch']['scope']['input_revision'] += 1
    response = client.post(url+'/commits', headers=headers, json=value)
    assert response.status_code == 409 and response.json()['code'] == 'VERSION_CONFLICT'
    assert client.get(url).json() == before
    assert client.get(url+'/receipts/'+value['request_id']).json() is None
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_profile_receipts')) == 1
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_model_attempts')) == 0
