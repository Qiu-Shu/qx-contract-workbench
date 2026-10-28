"""Saved material results become editable drafts without model or file mutation."""
import asyncio
from io import BytesIO
import json
from uuid import uuid4

from docx import Document
import pytest
from sqlalchemy import text

from cn_contract import domain as d
from cn_contract.assistant import material_attempts, service, attachment_files
from cn_contract.assistant.model_integrity import validate_model_records
from cn_contract.assistant.providers import attempts
from cn_contract.assistant.providers.policy import encoded
from cn_contract.assistant.runner import execute_attempt
from cn_contract.ops.storage import RestoreFenceRow
from cn_contract.security import provision_user, authenticate, revoke_session
from test_assistant_generation import world
from test_assistant_generation_routes import signin
from test_assistant_material_generation import material_world, selection_for, grant
from test_assistant_material_routes import boot
from test_assistant_main_provider import config
from test_contract_c4_material_flow import MaterialProvider


def saved(w, state='received'):
    database, actor, _, _ = w
    provider = MaterialProvider(config())
    if state not in ('received', 'reserved'):
        async def no_output(*args, **kwargs):
            return d.MainMaterialReply(state='received' if state == 'no_output' else state,
                error_code='MODEL_OUTPUT_INVALID' if state == 'no_output' else 'PROVIDER_ERROR')
        provider.generate_material = no_output
    selection = selection_for(w, 'contract')
    async def scenario():
        _, lease, basis = grant(w, selection, provider)
        with database.session() as db:
            result = material_attempts.enqueue(actor, selection, basis, db, lease)
        if state != 'reserved':
            result = await execute_attempt(actor, result.id, database, lease)
        return result
    return asyncio.run(scenario())


def path(http, attempt):
    return http.url+'/material-generations/'+str(attempt.id)+'/document'


def test_received_result_exports_editable_text_sources_and_unchanged_saved_data(material_world, tmp_path):
    result = saved(material_world)
    database = material_world[0]
    with database.session() as db:
        before = tuple(db.execute(text('SELECT * FROM assistant_model_outputs')).all())
        materials = tuple(db.execute(text('SELECT * FROM assistant_attachments')).all())
        original = db.execute(text('SELECT * FROM assistant_attachments')).mappings().one()
        original_bytes = attachment_files.read_checked(original, database)
    with boot(material_world, tmp_path) as http:
        response = http.client.get(path(http, result))
        assert response.status_code == 200, response.text
        assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        assert response.headers['content-disposition'] == 'attachment; filename="material-draft.docx"'
        assert response.headers['cache-control'] == 'no-store'
        assert response.headers['x-content-type-options'] == 'nosniff'
        document = Document(BytesIO(response.content))
        paragraphs = [item.text for item in document.paragraphs]
        body = '\n'.join(paragraphs)
        assert result.output.candidate.text in paragraphs
        assert '律师内部' in body and '待律师核对' in body and '未修改原合同' in body
        assert '材料目录' in body and '独立虚构借款.docx' in body
        assert '来源核对' in body
        assert '独立虚构材料：2026年1月3日甲向乙借款30000元，约定3月3日归还。' in body
        document.add_paragraph('律师离线编辑测试')
        edited = BytesIO(); document.save(edited)
        assert Document(BytesIO(edited.getvalue())).paragraphs[-1].text == '律师离线编辑测试'
    with database.session() as db:
        assert tuple(db.execute(text('SELECT * FROM assistant_model_outputs')).all()) == before
        assert tuple(db.execute(text('SELECT * FROM assistant_attachments')).all()) == materials
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_model_attempts')) == 1
    assert attachment_files.read_checked(original, database) == original_bytes


def test_document_rejects_other_conversation_user_and_unsigned_request(material_world, tmp_path):
    result = saved(material_world)
    database, actor, _, _ = material_world
    with database.session() as db:
        other = service.create_conversation(actor, d.AssistantConversationCreateInput(request_id=uuid4()), db)
        provision_user('other', 'fixture123', {d.Role.HANDLING_LAWYER}, db, local_workbench_password=True)
    with boot(material_world, tmp_path) as http:
        wrong = '/api/assistant/conversations/'+str(other.scope.conversation_id)+'/material-generations/'+str(result.id)+'/document'
        assert http.client.get(wrong).status_code == 403
        signin(http.client, 'other')
        assert http.client.get(path(http, result)).status_code == 403
        http.client.cookies.clear()
        assert http.client.get(path(http, result)).status_code == 401


@pytest.mark.parametrize('state', ['reserved', 'sent_unknown', 'failed_before_send', 'no_output'])
def test_unfinished_or_missing_result_has_no_document(material_world, tmp_path, state):
    result = saved(material_world, state)
    assert result.output is None
    with boot(material_world, tmp_path) as http:
        response = http.client.get(path(http, result))
        assert response.status_code == 409, response.text
        assert response.headers['content-type'].startswith('application/json')


def test_render_does_not_freeze_session_permission(material_world, tmp_path, monkeypatch):
    from cn_contract.web import assistant_material_generation as routes
    result = saved(material_world)
    render = routes.build
    with boot(material_world, tmp_path) as http:
        def revoke_after_render(output):
            content = render(output)
            with http.database.session() as db:
                actor = authenticate(http.client.cookies.get('qx_session'), db)
                revoke_session(actor.session_id, db)
            return content
        monkeypatch.setattr(routes, 'build', revoke_after_render)
        response = http.client.get(path(http, result))
        assert response.status_code == 401
        assert response.headers['content-type'].startswith('application/json')


def test_received_output_with_incomplete_run_cannot_be_exported(material_world, tmp_path):
    result = saved(material_world)
    with boot(material_world, tmp_path) as http:
        with http.database.session() as db:
            db.execute(text("UPDATE assistant_runs SET state='failed' WHERE id=:id"), {'id':str(result.run_id)})
        assert http.client.get(path(http, result)).status_code == 409


@pytest.mark.parametrize('state,error_code', [
    ('received', 'MODEL_OUTPUT_INVALID'),
    ('received', 'PROVIDER_ERROR'),
    ('received', 'BLOCKED_LIVE'),
    ('sent_unknown', 'PROVIDER_ERROR'),
    ('failed_before_send', 'PROVIDER_ERROR'),
])
def test_received_output_with_error_or_conflicting_receipt_is_not_exported(material_world, tmp_path, state, error_code):
    result = saved(material_world)
    with boot(material_world, tmp_path) as http:
        with http.database.session() as db:
            validate_model_records(db)
            raw = json.loads(db.scalar(text('SELECT receipt_json FROM assistant_model_receipts WHERE attempt_id=:id'), {'id':str(result.id)}))
            raw.update(state=state, error_code=error_code)
            if state != 'received':
                raw.update(usage=None, provider_request_id=None, cost_status='unknown', estimated_cost_usd=None)
            receipt = encoded(d.MainReceipt.model_validate(raw))
            trigger = db.scalar(text("SELECT sql FROM sqlite_master WHERE name='immutable_assistant_model_receipts_update'"))
            db.execute(text('DROP TRIGGER immutable_assistant_model_receipts_update'))
            db.execute(text('UPDATE assistant_model_receipts SET receipt_json=:raw WHERE attempt_id=:id'),
                {'id':str(result.id), 'raw':receipt})
            db.execute(text(trigger))
        with http.database.session() as db:
            with pytest.raises(ValueError, match='ASSISTANT_MODEL_INTEGRITY'):
                validate_model_records(db)
        response = http.client.get(path(http, result))
        assert response.status_code == 409, response.headers
        assert response.headers['content-type'].startswith('application/json')


@pytest.mark.parametrize('damage', ['wire_digest', 'receipt_cost', 'finished_event'])
@pytest.mark.parametrize('during_render', [False, True])
def test_document_rechecks_complete_ledger_before_and_after_render(material_world, tmp_path, monkeypatch, damage, during_render):
    from cn_contract.web import assistant_material_generation as routes
    result = saved(material_world)
    render = routes.build
    rendered = []
    with boot(material_world, tmp_path) as http:
        def corrupt():
            table = {'wire_digest':'assistant_model_attempts', 'receipt_cost':'assistant_model_receipts',
                'finished_event':'assistant_run_events'}[damage]
            with http.database.session() as db:
                triggers = db.execute(text("SELECT name,sql FROM sqlite_master WHERE type='trigger' AND tbl_name=:table"),
                    {'table':table}).all()
                for name, _ in triggers:
                    db.execute(text('DROP TRIGGER "' + name.replace('"', '""') + '"'))
                if damage == 'wire_digest':
                    db.execute(text('UPDATE assistant_model_attempts SET request_sha256=:sha WHERE id=:id'),
                        {'sha':'a'*64, 'id':str(result.id)})
                elif damage == 'receipt_cost':
                    raw = json.loads(db.scalar(text('SELECT receipt_json FROM assistant_model_receipts WHERE attempt_id=:id'),
                        {'id':str(result.id)}))
                    raw['estimated_cost_usd'] = '0'
                    db.execute(text('UPDATE assistant_model_receipts SET receipt_json=:raw WHERE attempt_id=:id'),
                        {'id':str(result.id), 'raw':encoded(d.MainReceipt.model_validate(raw))})
                else:
                    db.execute(text("DELETE FROM assistant_run_events WHERE run_id=:id AND event_type='model.finished'"),
                        {'id':str(result.run_id)})
                for _, sql in triggers:
                    db.execute(text(sql))
            with http.database.session() as db:
                with pytest.raises(ValueError, match='ASSISTANT_MODEL_INTEGRITY'):
                    validate_model_records(db)

        def render_checked(output):
            rendered.append(output.id)
            content = render(output)
            if during_render:
                corrupt()
            return content

        monkeypatch.setattr(routes, 'build', render_checked)
        if not during_render:
            corrupt()
        response = http.client.get(path(http, result))
        assert response.status_code == 409, response.headers
        assert response.headers['content-type'].startswith('application/json')
        assert rendered == ([result.output.id] if during_render else [])


def test_plain_text_attempt_is_not_a_material_document(world, tmp_path):
    database, actor, conversation, request, runtime = world
    with database.session() as db:
        pending = attempts.enqueue_generation(actor, conversation.scope.conversation_id, request, db, runtime)
    result = asyncio.run(execute_attempt(actor, pending.id, database, runtime))
    w = (database, actor, conversation.scope.conversation_id, None)
    with boot(w, tmp_path) as http:
        assert http.client.get(path(http, result)).status_code == 409


def test_restored_copy_can_download_only_exact_saved_document(material_world, tmp_path):
    result = saved(material_world)
    with material_world[0].session() as db:
        db.add(RestoreFenceRow(ref=str(uuid4()), archive_sha256='a'*64, manifest_sha256='b'*64, payload='{}'))
    with boot(material_world, tmp_path) as http:
        url = path(http, result)
        response = http.client.get(url)
        assert response.status_code == 200, response.text
        assert result.output.candidate.text in '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        assert http.client.post(url, headers=http.headers).status_code == 409
        assert http.client.get(url+'/extra').status_code == 409
        assert http.client.get(url.replace('/document', '/documents')).status_code == 409


@pytest.mark.parametrize('damage', ['audience', 'scope', 'manifest', 'citations', 'id'])
def test_corrupted_saved_result_never_becomes_a_document(material_world, tmp_path, damage):
    result = saved(material_world)
    with boot(material_world, tmp_path) as http:
        with http.database.session() as db:
            raw = json.loads(db.scalar(text('SELECT output_json FROM assistant_model_outputs WHERE attempt_id=:id'), {'id':str(result.id)}))
            if damage == 'audience': raw['audience'] = 'client'
            elif damage == 'scope': raw['scope']['conversation_id'] = str(uuid4())
            elif damage == 'manifest': raw['manifest']['materials'][0]['name'] = '伪造材料.docx'
            elif damage == 'citations': raw['candidate']['citations'] = ['0'*64]
            else: raw['id'] = str(uuid4())
            trigger = db.scalar(text("SELECT sql FROM sqlite_master WHERE name='immutable_assistant_model_outputs_update'"))
            db.execute(text('DROP TRIGGER immutable_assistant_model_outputs_update'))
            db.execute(text('UPDATE assistant_model_outputs SET output_json=:raw WHERE attempt_id=:id'),
                {'id':str(result.id), 'raw':json.dumps(raw, ensure_ascii=False)})
            db.execute(text(trigger))
        response = http.client.get(path(http, result))
        assert response.status_code == 409, response.text
        assert response.headers['content-type'].startswith('application/json')
