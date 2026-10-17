"""Synthetic native attachment HTTP: bounded input, ownership and exact reads."""
import asyncio
import hashlib
from io import BytesIO
import json
from threading import Event, Thread
from tempfile import SpooledTemporaryFile
from uuid import uuid4
from zipfile import ZipFile, ZIP_STORED

from docx import Document
from fastapi.testclient import TestClient
from starlette.requests import Request
import pytest
import httpx
from sqlalchemy import text

from test_assistant_routes import native, sign_in, create, serve

MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


@pytest.fixture
def attachments(native):
    app, root = native
    doc = Document(); doc.add_paragraph('独立虚构咨询材料：客户自述付款三万元，尚未核验。')
    stream = BytesIO(); doc.save(stream); raw = stream.getvalue()
    database = app.state.database
    database._assistant_synthetic_hashes = frozenset({hashlib.sha256(raw).hexdigest()})
    with TestClient(app) as client:
        headers = sign_in(client); conversation = create(client, headers)
        url = '/api/assistant/conversations/'+conversation['scope']['conversation_id']
        value = {'request_id':str(uuid4()), 'expected_scope':conversation['scope'],
                 'original_name':'合成咨询.docx', 'mime':MIME,
                 'policy':{'local_processing_accepted':True}}
        yield client, database, headers, url, value, raw, root


def upload(fixture, *, value=None, raw=None, fields=None, headers=None):
    client, _, default_headers, url, default_value, default_raw, _ = fixture
    value = default_value if value is None else value
    return client.post(url+'/attachments', headers=default_headers if headers is None else headers,
        files=[('file',(value['original_name'],default_raw if raw is None else raw,value['mime'])),
               ('metadata',(None,json.dumps(value,ensure_ascii=False),'application/json'))]+(fields or []))


def test_framework_multipart_stream_bound_and_cleanup_without_os_temp(monkeypatch):
    from cn_contract.web.assistant_attachments import multipart_input
    from cn_contract.security import ServiceError
    from starlette import formparsers
    spools = []
    def tracked(*args,**kwargs):
        file = SpooledTemporaryFile(*args,**kwargs); spools.append(file)
        def no_disk(): raise AssertionError('original must not spool to OS temp')
        file.rollover = no_disk
        return file
    monkeypatch.setattr(formparsers,'SpooledTemporaryFile',tracked)
    raw = b'x'*(1024**2+256)
    encoded = httpx.Request('POST','http://testserver/upload',files=[
        ('file',('合成.docx',raw,MIME)),('metadata',(None,'{}'))])
    body = encoded.read()
    def request():
        offset = 0
        async def receive():
            nonlocal offset
            chunk = body[offset:offset+4096]; offset += len(chunk)
            return {'type':'http.request','body':chunk,'more_body':offset<len(body)}
        return Request({'type':'http','headers':[(b'content-type',encoded.headers['content-type'].encode()),
                                               (b'content-length',b'1')]},receive)
    async def success():
        form = await multipart_input(request(),len(body))
        try:
            assert await form['file'].read() == raw and form['metadata'] == '{}'
        finally:
            await form.close()
    asyncio.run(success())
    assert len(spools) == 1 and spools[0].closed and not spools[0]._rolled
    with pytest.raises(ServiceError) as failure:
        asyncio.run(multipart_input(request(),len(body)-1))
    assert failure.value.status_code == 413
    assert len(spools) == 2 and all(file.closed and not file._rolled for file in spools)


def test_roundtrip_source_selection_and_receipts(attachments):
    client, database, headers, url, value, raw, _ = attachments
    response = upload(attachments)
    assert response.status_code == 201, response.text
    receipt = response.json(); aid = receipt['attachment']['id']; base = url+'/attachments/'+aid
    assert receipt['materials_revision'] == 1
    assert receipt['attachment']['original_name'] == value['original_name']
    assert 'storage_ref' not in response.text
    assert upload(attachments).json() == receipt
    assert client.get(url+'/attachments/by-request/'+value['request_id']).json() == receipt
    assert client.get(url+'/attachments/by-request/'+str(uuid4())).json() is None
    assert client.get(base).json() == receipt['attachment']
    original = client.get(base+'/original')
    assert original.content == raw
    assert original.headers['content-type'] == 'application/octet-stream'
    assert original.headers['content-disposition'] == f'attachment; filename="original-{aid}.docx"'
    scope = client.get(url+'/materials').json()['scope']
    parse_value = {'request_id':str(uuid4()),'expected_scope':scope}
    parsed = client.post(base+'/parses',headers=headers,json=parse_value)
    assert parsed.status_code == 200, parsed.text
    assert parsed.json()['parse']['reading_status'] == 'readable'
    assert parsed.json()['parse']['review_status'] == 'unverified'
    assert client.post(base+'/parses',headers=headers,json=parse_value).json() == parsed.json()
    assert client.get(base+'/parses/by-request/'+parse_value['request_id']).json() == parsed.json()
    assert client.get(base+'/parses/by-request/'+str(uuid4())).json() is None
    result = client.get(base+'/parses/1').json()
    anchor = result['result']['anchors'][0]
    excerpt = client.get(base+'/parses/1/anchors/'+anchor['id'],params={'quote_sha256':anchor['quote_sha256']})
    assert excerpt.status_code == 200 and excerpt.json()['anchor'] == anchor
    assert client.get(base+'/parses/1/anchors/'+anchor['id'],params={'quote_sha256':'f'*64}).status_code == 409
    current = client.get(url+'/materials').json()
    assert current['items'] == [{'attachment_id':aid,'parse_revision':1}]
    selection_key = str(uuid4())
    selection = client.put(url+'/materials',headers=headers,json={'request_id':selection_key,
        'expected_scope':current['scope'],'items':[]})
    assert selection.status_code == 200 and selection.json()['materials_revision'] == 3
    selected_receipt = client.get(url+'/materials/by-request/'+selection_key).json()
    assert selected_receipt['materials_revision'] == 3 and selected_receipt['items'] == []
    assert client.get(url+'/materials/by-request/'+str(uuid4())).json() is None
    assert client.get(url+'/materials').json()['items'] == []
    assert client.get(base+'/original').content == raw
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM matters')) == 0
        assert db.scalar(text('SELECT COUNT(*) FROM intake_cases')) == 0


def test_authentication_owner_csrf_and_origin_precede_body(attachments,monkeypatch):
    from cn_contract.web import assistant_attachments as transport
    client, _, headers, url, _, _, _ = attachments
    async def unread(*args,**kwargs):
        raise AssertionError('unauthorized body must not be consumed')
    monkeypatch.setattr(transport,'multipart_input',unread)
    monkeypatch.setattr(transport,'body_input',unread)
    assert client.post(url+'/attachments').status_code == 403
    assert client.post(url+'/attachments',headers={**headers,'Origin':'https://other.invalid'}).status_code == 403
    assert client.post(url+'/attachments',headers={**headers,'Sec-Fetch-Site':'same-site'}).status_code == 403
    other = sign_in(client,'另一合成律师')
    assert client.post(url+'/attachments',headers=other).status_code == 403
    assert client.put(url+'/materials',headers=other).status_code == 403
    assert client.post(url+'/attachments/'+str(uuid4())+'/parses',headers=other).status_code == 403
    client.cookies.clear()
    assert client.post(url+'/attachments').status_code == 401


def test_revocation_after_multipart_read_is_rechecked(attachments,monkeypatch):
    from cn_contract.web import assistant_attachments as transport
    client, database, _, url, _, _, _ = attachments
    original = transport.multipart_input
    async def revoke(*args,**kwargs):
        value = await original(*args,**kwargs)
        with database.session() as db: db.execute(text('UPDATE sessions SET revoked=1'))
        return value
    monkeypatch.setattr(transport,'multipart_input',revoke)
    assert upload(attachments).status_code == 401
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_attachments')) == 0


@pytest.mark.parametrize('kind',['extra','duplicate','name','mime','bad_json','duplicate_json','oversize_metadata'])
def test_multipart_metadata_is_strict(attachments,kind):
    client, database, headers, url, value, raw, _ = attachments
    metadata = json.dumps(value)
    name, mime = value['original_name'], MIME
    extras = []
    if kind == 'extra': extras = [('unrecognized',(None,'x'))]
    if kind == 'duplicate': extras = [('metadata',(None,metadata))]
    if kind == 'name': name = '另外一份.docx'
    if kind == 'mime': mime = 'application/octet-stream'
    if kind == 'bad_json': metadata = '{'
    if kind == 'duplicate_json': metadata = '{"request_id":"x","request_id":"y"}'
    if kind == 'oversize_metadata': metadata = 'x'*(64*1024+1)
    response = client.post(url+'/attachments',headers=headers,
        files=[('file',(name,raw,mime)),('metadata',(None,metadata))]+extras)
    assert response.status_code in (400,413,422), response.text
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_attachments')) == 0


def test_untrusted_length_and_file_type_are_bounded(attachments):
    client,database,headers,url,value,_,_ = attachments
    database.settings = database.settings.model_copy(update={'upload_max_file_mib':1})
    response = upload(attachments,raw=b'x'*(1024**2+64*1024+1),headers={**headers,'Content-Length':'1'})
    assert response.status_code == 413
    assert upload(attachments,raw=b'%PDF-1.0 fake').status_code == 400
    assert upload(attachments,raw=b'PK\x03\x04truncated').status_code == 400
    assert client.get(url+'/materials').json()['items'] == []


def test_multipart_spool_stays_in_memory_and_closes_on_failure(attachments,monkeypatch):
    from starlette import formparsers
    _,database,_,_,_,raw,_ = attachments
    stream = BytesIO(raw)
    with ZipFile(stream,'a') as package:
        package.writestr('customXml/synthetic.bin',b'x'*(1024**2),compress_type=ZIP_STORED)
    raw = stream.getvalue()
    database._assistant_synthetic_hashes = frozenset({hashlib.sha256(raw).hexdigest()})
    spools = []
    def tracked(*args,**kwargs):
        file = SpooledTemporaryFile(*args,**kwargs); spools.append(file)
        def no_disk(): raise AssertionError('original must not spool to OS temp')
        file.rollover = no_disk
        return file
    monkeypatch.setattr(formparsers,'SpooledTemporaryFile',tracked)
    assert upload(attachments,raw=raw).status_code == 201
    assert len(spools) == 1 and spools[0].closed and not spools[0]._rolled
    database.settings = database.settings.model_copy(update={'upload_max_file_mib':1})
    assert upload(attachments,raw=b'x'*(1024**2+64*1024+1)).status_code == 413
    assert all(file.closed and not file._rolled for file in spools)


def test_attachment_and_parse_reads_remain_owner_bound(attachments):
    client,_,_,url,_,_,_ = attachments
    aid = upload(attachments).json()['attachment']['id']; base = url+'/attachments/'+aid
    sign_in(client,'另一合成律师')
    for suffix in ('/materials','/attachments/'+aid,'/attachments/'+aid+'/original',
                   '/attachments/'+aid+'/parses/1','/attachments/by-request/'+str(uuid4()),
                   '/attachments/'+aid+'/parses/by-request/'+str(uuid4())):
        assert client.get(url+suffix).status_code == 403


def test_parse_and_selection_reauthenticate_after_json(attachments,monkeypatch):
    from cn_contract.web import assistant_attachments as transport
    client,database,headers,url,_,_,_ = attachments
    aid = upload(attachments).json()['attachment']['id']
    scope = client.get(url+'/materials').json()['scope']
    original = transport.body_input
    async def revoke(*args,**kwargs):
        result = await original(*args,**kwargs)
        with database.session() as db: db.execute(text('UPDATE sessions SET revoked=1'))
        return result
    monkeypatch.setattr(transport,'body_input',revoke)
    request = {'request_id':str(uuid4()),'expected_scope':scope}
    assert client.post(url+'/attachments/'+aid+'/parses',headers=headers,json=request).status_code == 401
    headers = sign_in(client)
    assert client.put(url+'/materials',headers=headers,json={**request,'items':[]}).status_code == 401
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_attachment_parses')) == 0
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_material_snapshots')) == 1


@pytest.mark.parametrize('phase',['upload','parse'])
def test_slow_input_and_parser_allow_another_http_write(native,monkeypatch,phase):
    from cn_contract.web import assistant_attachments as transport
    app,_ = native
    doc = Document(); doc.add_paragraph('并发验证使用的独立虚构资料。')
    stream = BytesIO(); doc.save(stream); raw = stream.getvalue()
    app.state.database._assistant_synthetic_hashes = frozenset({hashlib.sha256(raw).hexdigest()})
    entered, release = Event(), Event(); responses = []
    original_body, original_parse = transport.multipart_input, transport.service.parse_attachment
    async def observing_body(request,limit):
        if request.headers.get('X-Synthetic-Slow') == 'yes': entered.set()
        return await original_body(request,limit)
    def delayed_parse(*args,**kwargs):
        entered.set(); assert release.wait(8)
        return original_parse(*args,**kwargs)
    monkeypatch.setattr(transport,'multipart_input',observing_body)
    if phase == 'parse': monkeypatch.setattr(transport.service,'parse_attachment',delayed_parse)
    with serve(app) as address, httpx.Client(base_url=address,trust_env=False,timeout=3) as client:
        headers = sign_in(client); view = create(client,headers)
        url = '/api/assistant/conversations/'+view['scope']['conversation_id']
        metadata = {'request_id':str(uuid4()),'expected_scope':view['scope'],
                    'original_name':'并发合成.docx','mime':MIME,'policy':{'local_processing_accepted':True}}
        files = [('file',(metadata['original_name'],raw,MIME)),('metadata',(None,json.dumps(metadata)))]
        if phase == 'parse':
            received = client.post(url+'/attachments',headers=headers,files=files)
            assert received.status_code == 201, received.text
            aid = received.json()['attachment']['id']
            view = client.get(url).json()
        request = httpx.Request('POST',address+url+'/attachments',files=files)
        encoded = request.read()
        def chunks():
            yield encoded[:1]
            assert release.wait(8)
            yield encoded[1:]
        def pending():
            with httpx.Client(base_url=address,cookies=client.cookies,trust_env=False,timeout=10) as slow:
                if phase == 'upload':
                    responses.append(slow.post(url+'/attachments',content=chunks(),headers={**headers,
                        'Content-Type':request.headers['content-type'],'X-Synthetic-Slow':'yes'}))
                else:
                    responses.append(slow.post(url+'/attachments/'+aid+'/parses',headers=headers,
                        json={'request_id':str(uuid4()),'expected_scope':view['scope']}))
        thread = Thread(target=pending); thread.start()
        try:
            assert entered.wait(4)
            saved = client.put(url+'/draft',headers=headers,json={'expected_scope':view['scope'],
                'expected_composer_revision':view['composer_revision'],'content':'并发时仍可保存的原话草稿'})
            assert saved.status_code == 200, saved.text
        finally:
            release.set(); thread.join(12)
        assert not thread.is_alive() and len(responses) == 1
        assert responses[0].status_code == (201 if phase == 'upload' else 200), responses[0].text
        assert client.get(url).json()['draft_input'] == '并发时仍可保存的原话草稿'


def test_recovery_exact_attachment_reads_and_fenced_writes(attachments):
    from cn_contract.ops import storage
    from cn_contract.web.recovery import attach_recovery
    client,database,headers,url,value,raw,_ = attachments
    aid = upload(attachments).json()['attachment']['id']; base = url+'/attachments/'+aid
    parse_key = str(uuid4())
    parsed = client.post(base+'/parses',headers=headers,json={'request_id':parse_key,
        'expected_scope':client.get(url+'/materials').json()['scope']}).json()
    anchor = parsed['parse']['result']['anchors'][0]
    # Same middleware installation as native recovery, before the fresh app starts.
    from cn_contract.web.app import create_app
    from cn_contract.web.local_server import load_seed_policy
    with database.session() as db:
        db.add(storage.RestoreFenceRow(ref=str(uuid4()),archive_sha256='a'*64,manifest_sha256='b'*64,payload='{}'))
    app = create_app(database.settings,database=database,purpose='assisted_draft',
                     assisted_policy=load_seed_policy(),local_materials_enabled=True)
    attach_recovery(app,database)
    with TestClient(app) as recovered:
        h = sign_in(recovered)
        for suffix in ('/materials','/attachments/'+aid,'/attachments/'+aid+'/original',
                       '/attachments/'+aid+'/parses/1','/attachments/by-request/'+value['request_id'],
                       '/attachments/'+aid+'/parses/by-request/'+parse_key,'/materials/by-request/'+value['request_id']):
            assert recovered.get(url+suffix).status_code == 200, suffix
        assert recovered.get(base+'/parses/1/anchors/'+anchor['id'],params={'quote_sha256':anchor['quote_sha256']}).status_code == 200
        assert recovered.get(base+'/original').content == raw
        assert recovered.post(url+'/attachments',headers=h).status_code == 409
        assert recovered.post(base+'/parses',headers=h,json={}).status_code == 409
        assert recovered.put(url+'/materials',headers=h,json={}).status_code == 409
        for suffix in ('/attachments','/attachments/'+aid+'/unreviewed','/attachments/'+aid+'/parses/0',
                       '/attachments/'+aid+'/parses/1/unsafe'):
            assert recovered.get(url+suffix).status_code == 409
