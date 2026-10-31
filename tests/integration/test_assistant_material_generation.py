"""Independent fictional DOCX, native consent, ledger and loopback model transport."""
import asyncio
from datetime import datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import json
import pytest
from fastapi.testclient import TestClient
from docx import Document
from sqlalchemy import text
from cn_contract import domain as d
from cn_contract.assistant import attachments, business, service, material_preparation, material_attempts
from cn_contract.assistant.providers import attempts
from cn_contract.assistant.providers.registry import MainModelRegistry
from cn_contract.assistant.providers.main import DeepSeekProvider
from cn_contract.assistant.native_model_control import NativeModelControl
from cn_contract.assistant.runner import execute_attempt
from cn_contract.assistant.model_integrity import validate_model_records
from cn_contract.assistant.providers.policy import sha, encoded
from cn_contract.security import ServiceError, login, authenticate
from cn_contract.ops import backup, workspace_backup
from cn_contract.persistence import Database
from cn_contract.web import local_server
from tests.helpers.legacy_assistant_backup import PASSPHRASE
from test_assistant_generation import world
from test_assistant_main_provider import config, server, KEY, raw


@pytest.fixture
def material_world(world):
    database, _, conversation, _, _ = world
    with database.session() as db:
        token, _ = login('lawyer', 'fixture123', db)
        actor = authenticate(token, db)
    doc = Document(); doc.add_paragraph('独立虚构材料：2026年1月3日甲向乙借款30000元，约定3月3日归还。')
    doc.add_paragraph('乙称已经偿还10000元，但尚未提供转账凭证。')
    stream = BytesIO(); doc.save(stream); data = stream.getvalue()
    database._assistant_synthetic_hashes = frozenset({sha256(data).hexdigest()})
    cid = conversation.scope.conversation_id
    uploaded = attachments.ingest_attachment(actor, cid, d.AssistantAttachmentInput(request_id=uuid4(), expected_scope=conversation.scope,
        original_name='独立虚构借款.docx', mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        policy=d.AssistantMaterialPolicy(local_processing_accepted=True)), BytesIO(data), database=database)
    with database.session() as db:
        view = service.read_conversation(actor, cid, db)
    attachments.parse_attachment(actor, cid, uploaded.attachment.id,
        d.AssistantAttachmentParseInput(request_id=uuid4(), expected_scope=view.scope, baseline_view=None), database=database)
    return database, actor, cid, token


def selection_for(w, workflow='case'):
    database, actor, cid, _ = w
    with database.session() as db:
        view = service.read_conversation(actor, cid, db)
        parent = business.submit(actor, cid, d.BusinessInput(request_id=uuid4(), expected_scope=view.scope,
            expected_composer_revision=view.composer_revision, content='按原文核对虚构借款及还款争议。', workflow=workflow), db)
        view = service.read_conversation(actor, cid, db)
    return d.MainMaterialSelection(request_id=uuid4(), expected_scope=view.scope,
        expected_composer_revision=view.composer_revision, parent_business_run_id=parent.run_id)


def grant(w, selection, provider, budget='1'):
    database, actor, _, _ = w
    basis = material_preparation.preview(actor, selection, database, provider.config)
    registry = MainModelRegistry(); registry.configure(provider)
    policy = d.MainModelPolicy(policy_id=uuid4(), workspace_id=selection.expected_scope.workspace_id,
        actor_id=actor.user_id, config_sha256=sha(encoded(provider.config)), allowed_content_sha256s=(basis.manifest_sha256,),
        max_total_usd=budget, max_calls=1, expires_at=datetime.now(timezone.utc)+timedelta(minutes=5))
    registry.grant(policy, actor.session_id, basis=basis)
    return registry, registry.resolve(actor, selection.expected_scope), basis


def response(basis, citations=None):
    return raw({'id':'fictional-material-1', 'object':'chat.completion', 'created':1, 'model':'deepseek-flash', 'choices':[{'index':0,'finish_reason':'stop',
        'message':{'role':'assistant', 'content':json.dumps({'text':'借款与已还款金额存在争议，须核对原件及转账凭证。',
            'citations':citations or [basis.manifest.sources[0].source_id]}, ensure_ascii=False)}}],
        'usage':{'prompt_tokens':30,'completion_tokens':12,'total_tokens':42}})


@pytest.mark.parametrize('budget', ['1', None])
@pytest.mark.parametrize('workflow',['sales','case','contract','court'])
def test_all_four_workflows_send_exact_selected_materials_once_and_validate_archive(material_world, workflow, budget):
    w = material_world; database, actor, cid, _ = w
    selection = selection_for(w, workflow)
    basis = material_preparation.preview(actor, selection, database, config())
    assert len(basis.manifest.sources) == 3
    with server(response(basis)) as (url, received, _):
        async def scenario():
            registry, lease, approved = grant(w, selection, DeepSeekProvider.for_synthetic_loopback(config(), lambda:KEY, url), budget)
            with database.session() as db:
                first = material_attempts.enqueue(actor, selection, approved, db, lease)
            result = await execute_attempt(actor, first.id, database, lease)
            assert result.state == 'received' and result.output is not None
            assert result.output.manifest == approved.manifest
            assert result.output.review_status == 'unverified'
            assert await execute_attempt(actor, first.id, database, lease) == result
            with database.session() as db:
                assert material_attempts.enqueue(actor, selection, approved, db, lease) == result
                validate_model_records(db)
                from cn_contract.assistant.conversation_history import select_history
                assert not select_history(db, result.scope)
            return result
        asyncio.run(scenario())
        assert len(received) == 1
        sent = json.loads(received[0]['body'])
        assert json.loads(sent['messages'][1]['content']) == basis.manifest.model_dump(mode='json')


def test_concurrent_material_execution_claims_and_sends_only_once(material_world):
    w=material_world; database,actor,_,_=w; selection=selection_for(w)
    basis=material_preparation.preview(actor,selection,database,config())
    with server(response(basis),pause=.15) as (url,received,_):
        async def scenario():
            _,lease,approved=grant(w,selection,DeepSeekProvider.for_synthetic_loopback(config(),lambda:KEY,url))
            with database.session() as db:
                first=material_attempts.enqueue(actor,selection,approved,db,lease)
            results=await asyncio.gather(execute_attempt(actor,first.id,database,lease),
                execute_attempt(actor,first.id,database,lease))
            assert sorted(item.state for item in results)==['received','sending']
            completed=next(item for item in results if item.state=='received')
            assert await execute_attempt(actor,first.id,database,lease)==completed
            with database.session() as db:
                assert db.scalar(text('SELECT COUNT(*) FROM assistant_model_outputs'))==1
                validate_model_records(db)
        asyncio.run(scenario())
        assert len(received)==1


@pytest.mark.parametrize('budget', ['1', None])
def test_material_output_qxbak_roundtrip_is_read_only_and_never_resends(material_world, tmp_path, budget):
    w=material_world; database,actor,_,_=w; selection=selection_for(w)
    basis=material_preparation.preview(actor,selection,database,config())
    with server(response(basis)) as (url,received,_):
        async def scenario():
            _,lease,approved=grant(w,selection,DeepSeekProvider.for_synthetic_loopback(config(),lambda:KEY,url),budget)
            with database.session() as db:
                first=material_attempts.enqueue(actor,selection,approved,db,lease)
            return await execute_attempt(actor,first.id,database,lease)
        completed=asyncio.run(scenario())
        assert completed.output is not None and len(received)==1
    root=Path(database.engine.url.database).parent
    original=(root/'workbench.db').read_bytes()
    archive=tmp_path/'material-output.qxbak'
    workspace_backup.create_workspace_backup(root,archive,PASSPHRASE,data_parent=tmp_path)
    restored_info=workspace_backup.restore_workspace_backup(archive,tmp_path/'material-restored',PASSPHRASE)
    target=Path(restored_info['restored_root'])
    restored=Database('sqlite:///'+str(target/'workbench.db'),target/'protected',database.settings,initialize_schema=False)
    backup._mark_storage_flags(restored)
    try:
        with restored.session() as db:
            token,_=login('lawyer','fixture123',db)
            current=authenticate(token,db)
            assert attempts.read_attempt(current,completed.id,db)==completed
            validate_model_records(db)
    finally:
        restored.engine.dispose()
    app=local_server.build_local_app(target,data_parent=tmp_path,assistant_model_runtime=MainModelRegistry())
    with TestClient(app) as client:
        csrf=client.post('/api/login',json={'username':'lawyer','password':'fixture123'}).json()['csrf_token']
        headers={'X-CSRF-Token':csrf}
        url='/api/assistant/conversations/'+str(selection.expected_scope.conversation_id)
        for suffix in ('/material-preview','/material-generations'):
            blocked=client.post(url+suffix,headers=headers,json=selection.model_dump(mode='json'))
            assert blocked.status_code==409,blocked.text
    assert len(received)==1 and (root/'workbench.db').read_bytes()==original


def test_preview_is_readonly_and_overlimit_does_not_truncate(material_world):
    w=material_world; database, actor, _, _=w; selection=selection_for(w)
    with database.session() as db:
        before=db.scalar(text('SELECT COUNT(*) FROM assistant_runs'))
    with pytest.raises(ValueError):
        material_preparation.preview(actor, selection, database, config(max_input_bytes=100))
    with database.session() as db:
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_runs')) == before
        assert db.scalar(text('SELECT COUNT(*) FROM assistant_model_attempts')) == 0


@pytest.mark.parametrize('budget', ['1', None])
def test_native_material_consent_is_independent_from_local_policy_and_bound_to_current_selection(material_world, budget):
    w=material_world; database, actor, cid, token=w; selection=selection_for(w)
    async def scenario():
        registry=MainModelRegistry()
        control=NativeModelControl(database,registry)
        configured=control.configure(d.NativeModelConfigure(config=config(),secret=KEY))
        basis=material_preparation.preview(actor, selection, database, config())
        payload=d.NativeMaterialAuthorization(session_token=token,expected_config_version=configured['config_version'],
            basis=basis,max_total_usd=budget,valid_for_seconds=120)
        assert control.authorize_material(payload)=={'status':'authorized'}
        with database.session() as db:
            policy=db.scalar(text('SELECT policy_json FROM assistant_attachments LIMIT 1'))
            assert json.loads(policy)['external_processing'] is False
            view=service.read_conversation(actor,cid,db)
        attachments.select_materials(actor,cid,d.AssistantMaterialSelectionInput(request_id=uuid4(),expected_scope=view.scope,items=()),database=database)
        with pytest.raises(ServiceError):control.authorize_material(payload)
    asyncio.run(scenario())


def test_revoked_lease_prevents_send_and_records_zero_cost(material_world):
    w=material_world; database,actor,_,_=w; selection=selection_for(w)
    async def scenario():
        provider=DeepSeekProvider.for_synthetic_loopback(config(),lambda:pytest.fail('no key lookup'), 'http://127.0.0.1:1/chat/completions')
        registry,lease,basis=grant(w,selection,provider)
        with database.session() as db:first=material_attempts.enqueue(actor,selection,basis,db,lease)
        registry.revoke(actor.session_id)
        result=await execute_attempt(actor,first.id,database,lease)
        assert result.state=='failed_before_send' and result.receipt.cost_status=='not_sent'
        with database.session() as db:validate_model_records(db)
    asyncio.run(scenario())


def test_unknown_source_is_rejected_but_received_usage_is_preserved(material_world):
    w=material_world; database,actor,_,_=w; selection=selection_for(w)
    basis=material_preparation.preview(actor,selection,database,config())
    with server(response(basis,['a'*64])) as (url,received,_):
        async def scenario():
            _,lease,basis=grant(w,selection,DeepSeekProvider.for_synthetic_loopback(config(),lambda:KEY,url))
            with database.session() as db:first=material_attempts.enqueue(actor,selection,basis,db,lease)
            result=await execute_attempt(actor,first.id,database,lease)
            assert result.output is None and result.receipt.error_code=='MODEL_OUTPUT_INVALID'
            assert result.receipt.usage.total_tokens==42
            with database.session() as db:validate_model_records(db)
        asyncio.run(scenario())
        assert len(received)==1


@pytest.mark.parametrize('change', ['revoke', 'config_in_place', 'expired', 'materials'])
def test_transport_rechecks_material_permission_after_secret_resolution(material_world, change, monkeypatch):
    w=material_world; database,actor,cid,_=w; selection=selection_for(w)
    basis=material_preparation.preview(actor,selection,database,config())
    with server(response(basis)) as (url,received,_):
        async def scenario():
            def secret():
                if change == 'revoke':
                    registry.revoke(actor.session_id)
                elif change == 'config_in_place':
                    object.__setattr__(provider.config, 'timeout_seconds', 31)
                elif change == 'materials':
                    current=attachments.read_materials(actor,cid,database=database)
                    attachments.select_materials(actor,cid,d.AssistantMaterialSelectionInput(request_id=uuid4(),
                        expected_scope=current.scope,items=()),database=database)
                else:
                    from cn_contract.assistant.providers import registry as registry_module
                    class ExpiredDatetime:
                        @staticmethod
                        def now(tz=None):
                            return datetime.now(timezone.utc)+timedelta(hours=1)
                    monkeypatch.setattr(registry_module, 'datetime', ExpiredDatetime)
                return KEY
            provider=DeepSeekProvider.for_synthetic_loopback(config(),secret,url)
            registry,lease,approved=grant(w,selection,provider)
            with database.session() as db:
                first=material_attempts.enqueue(actor,selection,approved,db,lease)
            result=await execute_attempt(actor,first.id,database,lease)
            assert result.state=='failed_before_send'
            assert result.receipt.cost_status=='not_sent'
            with database.session() as db:validate_model_records(db)
        asyncio.run(scenario())
        assert received==[]
